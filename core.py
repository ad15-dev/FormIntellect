import io
import json
import os
import random
import re
import time
import traceback
import urllib.parse
import threading
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from collections import Counter

import pandas as pd
import requests
import pytz

from utils import IST, MIN_SECONDS_PER_ROW, STATUS_NAMES, BLANK, SUPPORTED_EXT, ALLOW_LOCAL_PATHS, fmt_clock, format_human, format_mm_ss

# --- Errors ---

@dataclass
class Problem:
    row: int | None
    column: str
    value: str
    problem: str
    fix: str

class AppError(Exception):
    def __init__(self, title, *, where="", why="", fix="", problems=None, technical=""):
        super().__init__(title)
        self.title = title
        self.where = where
        self.why = why
        self.fix = fix
        self.problems: list[Problem] = problems or []
        self.technical = technical

def explain_exception(exc):
    tech = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__))[-3000:]
    if isinstance(exc, PermissionError):
        return dict(title="Couldn't save the updated file", why="The file is open in another program.", fix="Close the file and try again.", technical=tech)
    if isinstance(exc, OSError) and getattr(exc, "errno", None) == 28:
        return dict(title="Server out of disk space", why="No room to save.", fix="Try again later.", technical=tech)
    return dict(title="Something unexpected went wrong", why="Unknown error.", fix="Contact support.", technical=tech)

# --- Data Loading ---

def sheet_csv_url(url):
    url = url.strip()
    if re.match(r"https://docs\.google\.com/spreadsheets/d/e/", url):
        if "output=csv" in url: return url
        base = url.split("?")[0].split("#")[0]
        if base.endswith("/pubhtml"): base = base[: -len("/pubhtml")] + "/pub"
        return base + "?output=csv"
    m = re.match(r"https://docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)", url)
    if not m: return None
    out = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    gid = re.search(r"[#&?]gid=(\d+)", url)
    if gid: out += f"&gid={gid.group(1)}"
    return out

def _read_file(path, name):
    lower = name.lower()
    if lower.endswith((".xlsx", ".xls")): return pd.read_excel(path)
    if lower.endswith(".ods"): return pd.read_excel(path, engine="odf")
    if lower.endswith(".tsv"): return pd.read_csv(path, sep="\t")
    if lower.endswith(".csv"): return pd.read_csv(path)
    if lower.endswith(".json"): return pd.read_json(path)
    raise AppError("Unsupported file type", where="Step 1", why=f"'{name}' is not supported.", fix="Use .xlsx or .csv.")

def load_dataframe(inp, workdir, log):
    kind = inp["source_kind"]
    if kind == "sheet":
        raw = inp["sheet_url"].strip()
        if not raw: raise AppError("No Google Sheet link", where="Step 1", why="Box is empty.", fix="Paste link.")
        url = sheet_csv_url(raw)
        if not url: raise AppError("Invalid Sheets link", where="Step 1", why="Not a valid Sheets URL.", fix="Copy from address bar.", technical=raw)
        try:
            resp = requests.get(url, timeout=30)
        except requests.exceptions.RequestException as exc:
            raise AppError("Connection failed", where="Step 1", why="Could not reach Google.", fix="Check internet.", technical=str(exc))
        if resp.status_code != 200 or resp.text.lstrip().startswith("<"):
            raise AppError("Sheet not public", where="Step 1", why="Google asked for login.", fix="Set sharing to 'Anyone with link'.", technical=f"HTTP {resp.status_code}")
        df = pd.read_csv(io.StringIO(resp.text))
        log(f"Google Sheet loaded: {len(df)} rows")
        return df, "Google Sheet", os.path.join(workdir, "responses_updated.xlsx")

    if kind == "manual":
        text = inp.get("manual_data", "").strip()
        if not text: raise AppError("No data entered", where="Step 1", why="Text box empty.", fix="Paste CSV data.")
        try:
            df = pd.read_csv(io.StringIO(text))
        except Exception as exc:
            raise AppError("Invalid CSV", where="Step 1", why="Could not parse text.", fix="Check format.", technical=str(exc))
        if df.empty: raise AppError("No rows", where="Step 1", why="Only headers found.", fix="Add data rows.")
        log(f"Manual data loaded: {len(df)} rows")
        return df, "Manual Entry", os.path.join(workdir, "manual_responses_updated.xlsx")

    if kind == "upload":
        if not inp.get("upload"): raise AppError("No file", where="Step 1", why="Nothing uploaded.", fix="Upload file.")
        name, data = inp["upload"]
        name = os.path.basename(name)
        path = os.path.join(workdir, name)
        with open(path, "wb") as fh: fh.write(data)
    else: # local
        path = inp["local_path"].strip().strip('"').strip("'")
        name = os.path.basename(path)
        if not path: raise AppError("No path", where="Step 1", why="Path empty.", fix="Enter path.")
        if not os.path.isfile(path): raise AppError("File not found", where="Step 1", why=f"No file at '{path}'.", fix="Check path.")

    try:
        df = _read_file(path, name)
    except AppError: raise
    except Exception as exc:
        raise AppError("Cannot open file", where=f"Step 1 · {name}", why="File may be damaged.", fix="Save as .xlsx/.csv.", technical=str(exc))

    stem, ext = os.path.splitext(name)
    if ext.lower() == ".xls": ext = ".xlsx"
    out_dir = workdir if kind == "upload" else os.path.dirname(os.path.abspath(path))
    return df, name, os.path.join(out_dir, f"{stem}_updated{ext}")

def save_dataframe(df, file_path):
    root, ext = os.path.splitext(file_path)
    tmp = f"{root}.saving{ext}"
    lower = ext.lower()
    if lower == ".xlsx": df.to_excel(tmp, index=False)
    elif lower == ".tsv": df.to_csv(tmp, sep="\t", index=False)
    elif lower == ".json": df.to_json(tmp, orient="records", indent=2)
    elif lower == ".ods": df.to_excel(tmp, index=False, engine="odf")
    else: df.to_csv(tmp, index=False)
    os.replace(tmp, file_path)

# --- Form Metadata ---

def extract_entry_ids(prefilled_link):
    parsed = urllib.parse.urlparse(prefilled_link)
    pairs = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    seen, ids = set(), []
    for key, _ in pairs:
        if key.startswith("entry.") and key not in seen:
            ids.append(key)
            seen.add(key)
    return ids

def _parse_fb_blob(html_text):
    token = "FB_PUBLIC_LOAD_DATA_"
    idx = html_text.find(token)
    if idx == -1: return None
    start = html_text.find("[", idx)
    if start == -1: return None
    depth, in_str, escape = 0, False, False
    for i in range(start, len(html_text)):
        ch = html_text[i]
        if in_str:
            if escape: escape = False
            elif ch == "\\": escape = True
            elif ch == '"': in_str = False
            continue
        if ch == '"': in_str = True
        elif ch == "[": depth += 1
        elif ch == "]":
            depth -= 1
            if depth == 0:
                try: return json.loads(html_text[start:i + 1])
                except json.JSONDecodeError: return None
    return None

@dataclass
class FormQuestion:
    entry: str
    label: str
    options: list
    required: bool
    grid_row: str = ""

def _questions_from_blob(blob):
    out = []
    try: items = blob[1][1]
    except (IndexError, TypeError): return out
    for item in items or []:
        if not isinstance(item, list) or len(item) < 5: continue
        label = item[1] if isinstance(item[1], str) else ""
        question = item[4]
        if not isinstance(question, list): continue
        is_grid = sum(1 for q in question if isinstance(q, list) and q and q[0] is not None) >= 2
        for q_item in question:
            if not isinstance(q_item, list) or not q_item or q_item[0] is None: continue
            opts = []
            if len(q_item) > 1 and isinstance(q_item[1], list):
                opts = [o[0] for o in q_item[1] if isinstance(o, list) and o and isinstance(o[0], str) and o[0]]
            required = len(q_item) > 2 and q_item[2] in (1, True)
            sub = ""
            if is_grid:
                lbl = q_item[3] if len(q_item) > 3 else None
                if isinstance(lbl, list) and lbl and isinstance(lbl[0], str): lbl = lbl[0]
                sub = lbl if isinstance(lbl, str) and lbl else f"row {q_item[0]}"
            out.append(FormQuestion(f"entry.{q_item[0]}", label, opts, required, sub))
    return out

@dataclass
class FormInfo:
    field_types: dict = field(default_factory=dict)
    options: dict = field(default_factory=dict)
    labels: dict = field(default_factory=dict)
    required: set = field(default_factory=set)
    reachable: bool = False

def fetch_form_metadata(prefilled_link):
    info = FormInfo()
    try: resp = requests.get(prefilled_link, timeout=30)
    except requests.exceptions.RequestException: return info
    if resp.status_code != 200: return info
    info.reachable = True
    page = resp.text
    for m in re.finditer(r'name="(entry\.\d+)"[^>]*type="([^"]+)"', page):
        eid, itype = m.group(1), m.group(2).lower()
        if eid not in info.field_types: info.field_types[eid] = itype
        elif "checkbox" in (info.field_types[eid], itype): info.field_types[eid] = "checkbox"
    for m in re.finditer(r'<textarea[^>]*name="(entry\.\d+)"', page):
        info.field_types.setdefault(m.group(1), "textarea")
    for m in re.finditer(r'<select[^>]*name="(entry\.\d+)"', page):
        info.field_types.setdefault(m.group(1), "select")
    blob = _parse_fb_blob(page)
    if blob:
        questions = _questions_from_blob(blob)
        for q in questions:
            if q.options: info.options[q.entry] = q.options
            info.labels[q.entry] = f"{q.label} — {q.grid_row}" if q.grid_row else q.label
            if q.required: info.required.add(q.entry)
            if q.grid_row: info.field_types[q.entry] = "grid"
    return info

# --- Validation & Normalization ---

def normalize_value(raw, entry_type="text"):
    if pd.isna(raw): return BLANK
    if isinstance(raw, (pd.Timestamp, datetime)):
        if entry_type == "date": return raw.strftime("%Y-%m-%d")
        if entry_type == "time": return raw.strftime("%H:%M")
        if entry_type == "datetime-local": return raw.strftime("%Y-%m-%dT%H:%M")
    if isinstance(raw, float) and raw.is_integer(): raw = int(raw)
    text = str(raw).strip()
    if not text: return BLANK
    if entry_type == "datetime-local" and " " in text and "T" not in text:
        text = text.replace(" ", "T", 1)
    return text

def is_blank(value): return value == BLANK

def best_option(value, options):
    if value in options: return value
    squash = lambda s: re.sub(r"\s+", " ", str(s)).strip().lower()
    by_key = {squash(o): o for o in options}
    if squash(value) in by_key: return by_key[squash(value)]
    try:
        num = float(value)
        for o in options:
            try:
                if float(o) == num: return o
            except (TypeError, ValueError): pass
    except ValueError: pass
    close = difflib.get_close_matches(squash(value), list(by_key), n=1, cutoff=0.85)
    return by_key[close[0]] if close else None

def split_choices(value, options):
    if value in options: return [value]
    return [v.strip() for v in value.split(",") if v.strip()]

def _opts_hint(opts, limit=8):
    shown = ", ".join(f"“{o}”" for o in opts[:limit])
    return shown + (f" (+{len(opts) - limit} more)" if len(opts) > limit else "")

def auto_correct(df, pending_indices, required_columns, form: FormInfo):
    corrections, problems = [], []
    for idx in pending_indices:
        for col, entry in required_columns.items():
            etype = form.field_types.get(entry, "text")
            value = normalize_value(df.at[idx, col], etype)
            opts = form.options.get(entry)
            new_val, reason = value, None
            if is_blank(value):
                if entry in form.required:
                    problems.append(Problem(sheet_row(idx), col, "(empty)", "Required question is empty.", f"Fill in an answer{': ' + _opts_hint(opts) if opts else ''}."))
                continue
            if opts:
                if etype == "checkbox":
                    parts = split_choices(value, opts)
                    fixed = [best_option(p, opts) for p in parts]
                    bad = [p for p, f in zip(parts, fixed) if f is None]
                    if bad:
                        problems.append(Problem(sheet_row(idx), col, value, f"{', '.join(f'“{b}”' for b in bad)} not in choices.", f"Use: {_opts_hint(opts)}."))
                        continue
                    candidate = ", ".join(dict.fromkeys(fixed))
                    if candidate != value: new_val, reason = candidate, "matched choices"
                else:
                    candidate = best_option(value, opts)
                    if candidate is None:
                        problems.append(Problem(sheet_row(idx), col, value, "Answer not in choices.", f"Use one of: {_opts_hint(opts)}."))
                        continue
                    if candidate != value: new_val, reason = candidate, "matched choice"
            if etype in ("date", "time", "datetime-local"):
                fmt, what, example = {"date": ("%Y-%m-%d", "date", "2026-10-29"), "time": ("%H:%M", "time", "14:30"), "datetime-local": ("%Y-%m-%dT%H:%M", "date/time", "2026-10-29 14:30")}[etype]
                try:
                    new_val = pd.to_datetime(value).strftime(fmt)
                    if new_val != value: reason = f"{what} rewritten"
                except (ValueError, TypeError, OverflowError):
                    problems.append(Problem(sheet_row(idx), col, value, f"Invalid {what}.", f"Write like {example}."))
                    continue
            if new_val != value:
                df.at[idx, col] = new_val
                corrections.append(dict(Row=sheet_row(idx), Column=col, From=value, To=new_val, Why=reason))
    return corrections, problems

def validate_data(df, pending_indices, required_columns, form: FormInfo):
    problems = []
    for idx in pending_indices:
        for col, entry in required_columns.items():
            etype = form.field_types.get(entry, "text")
            value = normalize_value(df.at[idx, col], etype)
            if is_blank(value): continue
            opts = form.options.get(entry)
            if opts:
                parts = split_choices(value, opts) if etype == "checkbox" else [value]
                bad = [p for p in parts if p not in opts]
                if bad:
                    problems.append(Problem(sheet_row(idx), col, value, "Answer not in choices.", f"Use: {_opts_hint(opts)}."))
                    continue
    return problems

def sheet_row(idx): return int(idx) + 2

# --- Scheduling ---

def base_window(start_t, end_t):
    today = datetime.now(IST).date()
    start_dt = IST.localize(datetime.combine(today, start_t))
    end_dt = IST.localize(datetime.combine(today, end_t))
    if end_dt <= start_dt: end_dt += timedelta(days=1)
    return start_dt, end_dt

def compute_time_window(start_t, end_t, start_range, end_range, total_rows, log):
    min_required_sec = total_rows * MIN_SECONDS_PER_ROW
    start_dt, end_dt = base_window(start_t, end_t)
    original_window_sec = int((end_dt - start_dt).total_seconds())
    srmin, srmax = sorted(start_range)
    ermin, ermax = sorted(end_range)
    rs = re_ = window_sec = 0
    rand_start = rand_end = None
    MAX_SCHEDULE_ATTEMPTS = 1000
    for _ in range(MAX_SCHEDULE_ATTEMPTS):
        rs = random.randint(srmin, srmax)
        re_ = random.randint(ermin, ermax)
        rand_start = start_dt + timedelta(seconds=rs)
        rand_end = end_dt + timedelta(seconds=re_)
        if rand_end <= rand_start: rand_end = rand_start + timedelta(seconds=1)
        window_sec = int((rand_end - rand_start).total_seconds())
        if window_sec >= original_window_sec and window_sec >= min_required_sec: break
    else:
        if original_window_sec < min_required_sec:
            raise AppError("Window too short", where="Step 4", why=f"{total_rows} rows need {format_human(min_required_sec)}.", fix="Increase window.")
        raise AppError("Random delays invalid", where="Step 4", why="Offsets make window shorter than entered.", fix="Adjust ranges.")
    
    log("========== TIME CONFIGURATION ==========")
    log(f"Entered start: {start_dt.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Entered end: {end_dt.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Actual start: {rand_start.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Actual end: {rand_end.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Rows: {total_rows}")
    log("========================================")
    
    cfg = dict(entered_start=start_dt, entered_end=end_dt, rs=rs, re=re_, rand_start=rand_start, rand_end=rand_end, original_window=original_window_sec, window=window_sec, min_required=min_required_sec, total_rows=total_rows)
    return rand_start, rand_end, window_sec, cfg

def wait_until(target_dt, stop):
    while True:
        remaining = (target_dt - datetime.now(IST)).total_seconds()
        if remaining <= 0: return True
        if stop.wait(min(remaining, 1.0)): return False

# --- Submission ---

@dataclass
class SubmitResult:
    ok: bool
    http: int
    problem: str = ""
    fix: str = ""
    technical: str = ""

def build_payload(row, required_columns, form: FormInfo):
    payload = {}
    for col, entry in required_columns.items():
        etype = form.field_types.get(entry, "text")
        value = normalize_value(row[col], etype)
        opts = form.options.get(entry)
        if is_blank(value) and opts: continue
        if etype == "checkbox": payload[entry] = split_choices(value, opts or [])
        else: payload[entry] = value
    return payload

def explain_response(status, text, url):
    snippet = text[:300] if text else ""
    tech = f"HTTP {status} · {url}\n{snippet}"
    low = (text or "").lower()
    if status == 200:
        if "accounts.google.com" in url or "servicelogin" in low:
            return SubmitResult(False, status, "Form requires sign-in.", "Turn off 'Restrict to users'.", tech)
        if "no longer accepting responses" in low:
            return SubmitResult(False, status, "Form closed.", "Enable 'Accepting responses'.", tech)
        return SubmitResult(True, status)
    if status == 0: return SubmitResult(False, status, "Connection failed.", "Check internet.", tech)
    if status == 400: return SubmitResult(False, status, "Google rejected answers.", "Check required fields/choices.", tech)
    if status in (401, 403): return SubmitResult(False, status, "Access denied.", "Check form permissions.", tech)
    if status == 404: return SubmitResult(False, status, "Form not found.", "Check link.", tech)
    if status == 429: return SubmitResult(False, status, "Too many requests.", "Increase time window.", tech)
    if status >= 500: return SubmitResult(False, status, "Google server error.", "Try again later.", tech)
    return SubmitResult(False, status, f"Unexpected error ({status}).", "Try again.", tech)

def submit_row(session, row, prep, stop, max_retries=3):
    payload = build_payload(row, prep.required_columns, prep.form)
    headers = {"User-Agent": "Mozilla/5.0", "Referer": prep.prefilled_link, "Origin": "https://docs.google.com"}
    last = SubmitResult(False, 0, technical="no attempt")
    for attempt in range(max_retries):
        try: r = session.post(prep.form_url, data=payload, headers=headers, timeout=30)
        except requests.exceptions.RequestException as exc:
            last = explain_response(0, "", prep.form_url)
            last.technical = f"{type(exc).__name__}: {exc}"
            delay = 2 * (2 ** attempt)
        else:
            if r.status_code != 429 and r.status_code < 500: return explain_response(r.status_code, r.text, r.url)
            last = explain_response(r.status_code, r.text, r.url)
            try: delay = float(r.headers.get("Retry-After", ""))
            except ValueError: delay = 2 * (2 ** attempt)
        if attempt < max_retries - 1 and stop.wait(min(delay, 60)): break
    last.technical = f"Tried {max_retries} times.\n{last.technical}"
    return last

# --- Analysis ---

def generate_analysis_sheet(df, file_path, submitted_count, log):
    try:
        from openpyxl import load_workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("Analysis skipped: openpyxl missing.")
        return
    if not file_path.lower().endswith(".xlsx"):
        log("Analysis skipped: not .xlsx.")
        return
    
    wb = load_workbook(file_path)
    if "Analysis" in wb.sheetnames: del wb["Analysis"]
    ws = wb.create_sheet("Analysis")
    NAVY, BLUE, LIGHT_BLUE, WHITE = "1F3864", "2E75B6", "DEEAF1", "FFFFFF"
    def solid(color): return PatternFill("solid", fgColor=color)
    thin = Side(style="thin", color="CCCCCC")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws.merge_cells("A1:F1")
    ws["A1"] = "Response Analysis"
    ws["A1"].fill = solid(NAVY); ws["A1"].font = Font(bold=True, color=WHITE, size=14); ws["A1"].alignment = ctr
    ws.row_dimensions[1].height = 28
    ws["A2"] = f"Generated: {datetime.now(IST).strftime('%d-%m-%Y %I:%M %p')}"
    ws["B2"] = f"Total rows: {len(df)}"
    ws["C2"] = f"Submitted: {submitted_count}"
    
    analysis_cols = [c for c in df.columns if str(c).strip().lower() not in STATUS_NAMES]
    current_row = 4
    for col_name in analysis_cols:
        series = df[col_name].dropna().astype(str).str.strip()
        series = series[series != ""]
        if series.empty: continue
        is_checkbox = series.str.contains(",").any()
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        cell = ws.cell(row=current_row, column=1, value=str(col_name))
        cell.fill = solid(BLUE); cell.font = Font(bold=True, color=WHITE, size=11); cell.alignment = ctr
        current_row += 1
        hdr_row = current_row
        for ci, h in enumerate(["#", "Answer", "Count", "% of Total"], 1):
            c = ws.cell(row=current_row, column=ci, value=h)
            c.fill = solid(BLUE); c.font = Font(bold=True, color=WHITE, size=10); c.alignment = ctr; c.border = bdr
        current_row += 1
        if is_checkbox:
            all_vals = []
            for v in series: all_vals.extend(x.strip() for x in v.split(",") if x.strip())
            freq = Counter(all_vals)
        else: freq = Counter(series.tolist())
        total_responders = len(series)
        items_sorted = sorted(freq.items(), key=lambda x: -x[1])
        data_start = current_row
        for ri, (answer, count) in enumerate(items_sorted, 1):
            pct = (count / total_responders * 100) if total_responders else 0
            for ci, val in enumerate([ri, answer, count, f"{pct:.1f}%"], 1):
                c = ws.cell(row=current_row, column=ci, value=val)
                c.fill = solid(LIGHT_BLUE if ri % 2 == 0 else WHITE); c.font = Font(color="1F1F1F", size=10); c.alignment = ctr; c.border = bdr
            current_row += 1
        total_row_idx = current_row
        total_pct = "100%" if not is_checkbox else "---"
        for ci, val in enumerate(["", "TOTAL", sum(freq.values()), total_pct], 1):
            c = ws.cell(row=current_row, column=ci, value=val)
            c.fill = solid(NAVY); c.font = Font(bold=True, color=WHITE, size=10); c.alignment = ctr; c.border = bdr
        current_row += 1
        if items_sorted:
            chart = BarChart(); chart.type = "col"; chart.grouping = "clustered"
            chart.title = str(col_name)[:30]; chart.y_axis.title = "Count"; chart.x_axis.title = "Answer"
            chart.style = 10; chart.width = 18; chart.height = 12
            data_ref = Reference(ws, min_col=3, max_col=3, min_row=hdr_row, max_row=total_row_idx - 1)
            cats_ref = Reference(ws, min_col=2, max_col=2, min_row=data_start, max_row=total_row_idx - 1)
            chart.add_data(data_ref, titles_from_data=True); chart.set_categories(cats_ref)
            ws.add_chart(chart, f"{get_column_letter(7)}{data_start}")
        current_row += 2
    wb.save(file_path)
    log("Analysis sheet added.")

# --- Preparation ---

@dataclass
class Prepared:
    df: pd.DataFrame
    output_path: str
    status_col: str
    required_columns: dict
    form: FormInfo
    pending_indices: list
    form_url: str
    prefilled_link: str
    rand_start: datetime
    rand_end: datetime
    window_sec: int
    pre_lines: list
    schedule_lines: list
    inp: dict
    summary: dict
    mapping: list
    corrections: list
    time_cfg: dict

def check_link(link):
    link = link.strip()
    how = "In Google Forms: ⋮ menu → 'Get pre-filled link' → Fill answers → 'Get link' → Copy."
    if not link: raise AppError("No form link", where="Step 2", why="Box empty.", fix=how)
    if "/viewform" not in link or not link.startswith("http"): raise AppError("Invalid link", where="Step 2", why="Not a pre-filled link.", fix=how, technical=link)
    if not extract_entry_ids(link): raise AppError("No entry IDs", where="Step 2", why="Link missing 'entry.' values.", fix=how, technical=link)
    return link

def map_columns(df, prefilled_link, status_col):
    entry_cols = [c for c in df.columns if str(c).startswith("entry.")]
    prefilled_ids = extract_entry_ids(prefilled_link)
    non_status_cols = [c for c in df.columns if c != status_col]
    if entry_cols: return {c: c for c in entry_cols}
    if prefilled_ids:
        matched = [c for c in prefilled_ids if c in df.columns]
        if matched: return {c: c for c in matched}
        if len(prefilled_ids) == len(non_status_cols): return dict(zip(non_status_cols, prefilled_ids))
    raise AppError("Columns mismatch", where="Step 1 & 2", why=f"Sheet has {len(non_status_cols)} cols, form has {len(prefilled_ids)} questions.", fix="Match counts or rename columns to entry.IDs.", technical=f"Cols: {list(map(str, non_status_cols))}\nIDs: {prefilled_ids}")

def prepare(inp: dict, workdir: str) -> Prepared:
    lines: list = []
    log = lines.append
    prefilled_link = check_link(inp["link"])
    form_url = prefilled_link.split("/viewform")[0] + "/formResponse"
    df, source_label, output_path = load_dataframe(inp, workdir, log)
    if df.empty: raise AppError("No rows", where="Step 1", why="Only headers found.", fix="Add data.")
    df.columns = [str(c).strip() for c in df.columns]
    n_rows, n_cols = len(df), len(df.columns)
    status_col = next((c for c in df.columns if c.lower() in STATUS_NAMES), None)
    if status_col is None:
        status_col = "Status"
        df[status_col] = ""
    required_columns = map_columns(df, prefilled_link, status_col)
    status_series = df[status_col].fillna("").astype(str).str.strip().str.lower()
    pending_indices = df[status_series != "submitted"].index.tolist()
    pending_total = len(pending_indices)
    if pending_total == 0: raise AppError("All submitted", where="Step 1", why="No pending rows.", fix=f"Clear '{status_col}' column.")
    if inp["random"]:
        random.shuffle(pending_indices)
        n = int(inp["sample_n"] or 0)
        if n > 0: pending_indices = pending_indices[: min(n, len(pending_indices))]
    total_rows = len(pending_indices)
    log(f"Form link: {prefilled_link}")
    log("Column → question:")
    for col, entry in required_columns.items(): log(f"  {col} → {entry}")
    form = fetch_form_metadata(prefilled_link)
    if form.field_types:
        log("Question types read:")
        for entry in required_columns.values(): log(f"  {entry}: {form.field_types.get(entry, 'unknown')}")
    else: log("Could not read question types.")
    corrections, problems = auto_correct(df, pending_indices, required_columns, form)
    if corrections:
        log(f"Auto-fixed {len(corrections)} values:")
        for c in corrections: log(f"  Row {c['Row']} | {c['Column']} | {c['From']} → {c['To']} ({c['Why']})")
    if not problems: problems = validate_data(df, pending_indices, required_columns, form)
    if problems:
        rows = sorted({p.row for p in problems})
        raise AppError(f"{len(problems)} cells need fixing", where=f"Sheet rows {', '.join(map(str, rows[:10]))}", why="Answers don't match form.", fix="Fix cells and upload again.", problems=problems)
    if corrections: save_dataframe(df, output_path)
    sched_lines: list = []
    rand_start, rand_end, window_sec, time_cfg = compute_time_window(inp["start"], inp["end"], inp["srange"], inp["erange"], total_rows, sched_lines.append)
    summary = dict(source=source_label, is_google_sheets=inp["source_kind"] == "sheet", rows=n_rows, cols=n_cols, pending=pending_total, selected=total_rows, random=inp["random"], entry_fields=len(extract_entry_ids(prefilled_link)), validated_fields=len(required_columns), form_reachable=form.reachable)
    mapping = [dict(Column=col, Question=form.labels.get(entry) or "—", ID=entry, Type=form.field_types.get(entry, "unknown"), Required="Yes" if entry in form.required else "") for col, entry in required_columns.items()]
    return Prepared(df=df, output_path=output_path, status_col=status_col, required_columns=required_columns, form=form, pending_indices=pending_indices, form_url=form_url, prefilled_link=prefilled_link, rand_start=rand_start, rand_end=rand_end, window_sec=window_sec, pre_lines=lines, schedule_lines=sched_lines, inp=inp, summary=summary, mapping=mapping, corrections=corrections, time_cfg=time_cfg)

def reroll_schedule(prep: Prepared):
    lines: list = []
    i = prep.inp
    rs_, re_, ws, cfg = compute_time_window(i["start"], i["end"], i["srange"], i["erange"], prep.summary["selected"], lines.append)
    prep.rand_start, prep.rand_end, prep.window_sec = rs_, re_, ws
    prep.time_cfg, prep.schedule_lines = cfg, lines

def console_text(prep: Prepared) -> str:
    return "\n".join(prep.pre_lines + [""] + prep.schedule_lines).strip()

# --- Background Job ---

class Job:
    def __init__(self, prep: Prepared):
        self.prep = prep
        self.total = len(prep.pending_indices)
        self.lock = threading.Lock()
        self.stop = threading.Event()
        self.status = "waiting"
        self.rows: list = []
        self.failures: list = []
        self.submitted = 0
        self.failed = 0
        self.next_target = None
        self.error = None
        self.notified = False
        self.t_start = None
        self.t_end = None
        self.thread = threading.Thread(target=self._run, daemon=True)

    def start(self): self.thread.start()
    @property
    def active(self): return self.status in ("waiting", "running")
    def _set(self, **kw):
        with self.lock:
            for k, v in kw.items(): setattr(self, k, v)
    def _finish(self, status): self._set(t_end=time.time(), status=status, next_target=None)
    def elapsed(self):
        if not self.t_start: return 0
        return int((self.t_end or time.time()) - self.t_start)
    def _add(self, **row):
        with self.lock: self.rows.append(row)
    def _note(self, text): self._add(kind="note", text=text)
    def snapshot(self):
        with self.lock:
            return {"status": self.status, "rows": list(self.rows), "failures": list(self.failures), "submitted": self.submitted, "failed": self.failed, "total": self.total, "next_target": self.next_target, "error": self.error}

    def _run(self):
        p = self.prep
        session = requests.Session()
        dirty, last_save = False, time.time()
        try:
            if datetime.now(IST) < p.rand_start:
                self._set(next_target=p.rand_start)
                self._note(f"Waiting for start time ({fmt_clock(p.rand_start)} IST)")
                if not wait_until(p.rand_start, self.stop):
                    self._note("Stopped before start.")
                    self._finish("stopped"); return
            self._set(t_start=time.time(), status="running")
            total_rows = self.total
            window_sec = max(p.window_sec, total_rows + 5)
            offsets = sorted(random.sample(range(1, window_sec + 1), total_rows))
            start_ref = datetime.now(IST)
            for i, (offset, idx) in enumerate(zip(offsets, p.pending_indices), 1):
                if self.stop.is_set(): break
                target = start_ref + timedelta(seconds=offset)
                self._set(next_target=target)
                if not wait_until(target, self.stop): break
                res = submit_row(session, p.df.loc[idx], p, self.stop)
                now_s = fmt_clock(datetime.now(IST))
                row_no = sheet_row(idx)
                if res.ok:
                    p.df.at[idx, p.status_col] = "submitted"
                    dirty = True
                    with self.lock: self.submitted += 1
                    if i < total_rows:
                        nxt = start_ref + timedelta(seconds=offsets[i])
                        gap = format_human(offsets[i] - offsets[i - 1])
                        self._add(kind="ok", row=row_no, time=now_s, nxt=fmt_clock(nxt), gap=gap, text=f"Sheet row {row_no} sent at {now_s} IST · next at {fmt_clock(nxt)} (in {gap})")
                    else: self._add(kind="final", row=row_no, time=now_s, text=f"Sheet row {row_no} sent at {now_s} IST · last entry")
                else:
                    fail = dict(row=row_no, time=now_s, problem=res.problem, fix=res.fix, technical=res.technical, http=res.http)
                    with self.lock: self.failed += 1; self.failures.append(fail)
                    self._add(kind="fail", **fail, text=f"Sheet row {row_no} NOT sent at {now_s} IST · {res.problem} · {res.fix}")
                if dirty and (time.time() - last_save >= 3 or i == total_rows):
                    save_dataframe(p.df, p.output_path)
                    dirty, last_save = False, time.time()
            if dirty: save_dataframe(p.df, p.output_path)
            if self.stop.is_set():
                self._note("Stopped by user.")
                self._finish("stopped"); return
            self._note("All rows processed.")
            try: generate_analysis_sheet(p.df, p.output_path, self.submitted, self._note)
            except Exception as exc: self._note(f"Analysis sheet failed ({type(exc).__name__}).")
            self._finish("done")
        except Exception as exc:
            if dirty:
                try: save_dataframe(p.df, p.output_path)
                except Exception: pass
            err = explain_exception(exc)
            self._set(error=err)
            self._note(f"ERROR: {err['title']}")
            self._finish("error")
        finally: session.close()

    def log_text(self):
        lines = list(self.prep.pre_lines) + [""] + list(self.prep.schedule_lines) + [""]
        with self.lock:
            lines += [r["text"] for r in self.rows]
            if self.failures:
                lines += ["", "========== FAILED ROWS =========="]
                for f in self.failures: lines += [f"Sheet row {f['row']}: {f['problem']}", f"  Fix: {f['fix']}", f"  Tech: {f['technical']}"]
            if self.error: lines += ["", "========== ERROR ==========", self.error["title"], self.error["technical"]]
        return "\n".join(lines)

    def failures_csv(self):
        with self.lock:
            rows = [dict(**{"Sheet row": f["row"], "Time (IST)": f["time"], "What went wrong": f["problem"], "What to do": f["fix"], "Technical detail": f["technical"]}) for f in self.failures]
        return pd.DataFrame(rows).to_csv(index=False)