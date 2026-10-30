"""
FormIntellect — Professional Google Form submission scheduler.
Single-file version with custom theme, execution table, and forced 12h UI.
"""
from __future__ import annotations

import difflib
import html
import io
import json
import os
import random
import re
import tempfile
import threading
import time
import traceback
import urllib.parse
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, time as dtime, timedelta

import pandas as pd
import pytz
import requests
import streamlit as st

try:
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

IST = pytz.timezone("Asia/Kolkata")
MIN_SECONDS_PER_ROW = 2
MAX_SCHEDULE_ATTEMPTS = 1000
STATUS_NAMES = {"status", "submit", "submitted"}
BLANK = "NA"
SAVE_EVERY_SEC = 3
ALLOW_LOCAL_PATHS = os.environ.get("FORMINTELLECT_ALLOW_LOCAL_PATHS") == "1"
SUPPORTED_EXT = (".xlsx", ".xls", ".csv", ".tsv", ".ods", ".json")

# ============================================================
#  CUSTOM THEME CSS
# ============================================================

CSS = """
:root {
  --bg-primary: #121212;
  --text-primary: #e2e8f0;
  --primary: #008a50;
  --primary-fg: #dde8e3;
  --secondary: #242424;
  --secondary-fg: #fafafa;
  --accent: #313131;
  --accent-fg: #fafafa;
  --card: #171717;
  --card-fg: #e2e8f0;
  --popover: #242424;
  --popover-fg: #a9a9a9;
  --muted: #1f1f1f;
  --muted-fg: #a2a2a2;
  --border: #333333;
  --radius: 8px;
}

* { box-sizing: border-box; }
html, body, .stApp { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: var(--bg-primary); color: var(--text-primary); }
.stApp { background: var(--bg-primary); }
.stApp p, .stApp li, .stApp label, .stMarkdown { color: var(--text-primary); }
header[data-testid="stHeader"] { background: transparent; }
#MainMenu, footer { visibility: hidden; }
.block-container { max-width: 960px; padding: 2rem 1rem; }
[data-testid="stCaptionContainer"], .stCaption { color: var(--muted-fg) !important; font-size: 0.85rem; }

.fi-hero { display: flex; align-items: center; gap: 1rem; margin-bottom: 0.5rem; padding: 1.5rem; background: var(--card); border-radius: var(--radius); border: 1px solid var(--border); }
.fi-logo { width: 48px; height: 48px; border-radius: var(--radius); background: var(--primary); display: flex; align-items: center; justify-content: center; color: var(--primary-fg); font-size: 1.5rem; font-weight: 700; }
.fi-name { font-size: 1.75rem; font-weight: 700; letter-spacing: -0.02em; line-height: 1.2; }
.fi-name span { color: var(--primary); }
.fi-tag { color: var(--muted-fg); font-size: 0.95rem; margin: 0 0 2rem 4rem; }

.fi-stepper { display: flex; align-items: center; gap: 1rem; margin: 0 0 2rem; padding: 1.25rem; background: var(--card); border: 1px solid var(--border); border-radius: var(--radius); }
.fi-st { display: flex; align-items: center; gap: 0.75rem; color: var(--muted-fg); font-size: 0.875rem; font-weight: 500; white-space: nowrap; }
.fi-st .n { width: 32px; height: 32px; border-radius: 50%; border: 2px solid var(--border); display: flex; align-items: center; justify-content: center; font-weight: 600; font-size: 0.875rem; }
.fi-st .l small { display: block; color: var(--muted-fg); font-size: 0.75rem; font-weight: 400; margin-top: 2px; }
.fi-st.active { color: var(--text-primary); }
.fi-st.active .n { border-color: var(--primary); color: var(--primary); background: rgba(0,138,80,0.1); }
.fi-st.done { color: var(--text-primary); }
.fi-st.done .n { background: var(--primary); border-color: var(--primary); color: var(--primary-fg); }
.fi-st-line { flex: 1; height: 2px; background: var(--border); min-width: 20px; }
.fi-st-line.done { background: var(--primary); }

.fi-sec h2 { font-size: 1.5rem; font-weight: 700; margin: 0 0 0.5rem; }
.fi-sec p { color: var(--muted-fg); font-size: 0.95rem; margin: 0 0 1.5rem; }

div[data-testid="stVerticalBlockBorderWrapper"] { background: var(--card) !important; border: 1px solid var(--border) !important; border-radius: var(--radius) !important; padding: 1.5rem; margin-bottom: 1.5rem; }
.fi-step { display: flex; align-items: center; gap: 0.75rem; margin-bottom: 0.5rem; }
.fi-step b { font-size: 0.75rem; font-weight: 700; color: var(--primary); background: rgba(0,138,80,0.1); border: 1px solid rgba(0,138,80,0.3); border-radius: 4px; padding: 0.25rem 0.625rem; }
.fi-step .t { font-weight: 600; font-size: 1.125rem; }
.fi-step-d { color: var(--muted-fg); font-size: 0.875rem; margin: 0 0 1rem 3rem; }

[data-testid="stWidgetLabel"] p { color: var(--muted-fg) !important; font-size: 0.875rem !important; font-weight: 500; }
div[data-baseweb="input"], div[data-baseweb="base-input"], div[data-baseweb="select"] > div, div[data-baseweb="textarea"] { background: var(--accent) !important; border-color: var(--border) !important; border-radius: var(--radius) !important; }
div[data-baseweb="input"]:focus-within { border-color: var(--primary) !important; }
input, textarea, select { color: var(--text-primary) !important; font-size: 0.95rem; }

/* Force 12h Time Picker Layout */
.time-picker-row { display: flex; gap: 0.5rem; align-items: center; }
.time-picker-row > div { flex: 1; min-width: 0; }
.time-picker-row [data-testid="stNumberInput"] input, 
.time-picker-row [data-testid="stSelectbox"] div[role="combobox"] { 
  text-align: center; font-family: monospace; font-size: 1.1rem; padding: 0.5rem; 
}

[data-testid="stFileUploaderDropzone"] { background: var(--accent) !important; border: 2px dashed var(--border) !important; border-radius: var(--radius) !important; padding: 2rem !important; transition: all 0.2s ease; }
[data-testid="stFileUploaderDropzone"]:hover { border-color: var(--primary) !important; background: rgba(0,138,80,0.05) !important; }
[data-testid="stFileUploaderDropzone"] p { color: var(--muted-fg) !important; font-size: 0.95rem !important; }

[data-testid="stExpander"] { background: var(--accent) !important; border: 1px solid var(--border) !important; border-radius: var(--radius) !important; margin: 1rem 0 !important; }
[data-testid="stExpander"] summary { padding: 1rem 1.25rem !important; cursor: pointer !important; }
[data-testid="stExpander"] summary p { color: var(--text-primary) !important; font-weight: 600 !important; font-size: 0.95rem !important; margin: 0 !important; }
[data-testid="stExpander"] .streamlit-expanderContent { background: var(--card) !important; padding: 1.25rem !important; border-top: 1px solid var(--border) !important; }
[data-testid="stExpander"] .streamlit-expanderContent p { color: var(--muted-fg) !important; line-height: 1.7 !important; margin-bottom: 0.75rem !important; }
[data-testid="stExpander"] .streamlit-expanderContent strong { color: var(--primary) !important; font-weight: 600 !important; }

.stButton > button, .stDownloadButton > button { border-radius: var(--radius); border: 1px solid var(--border); background: var(--accent); color: var(--accent-fg); font-weight: 600; padding: 0.625rem 1.25rem; transition: all 0.2s ease; font-size: 0.95rem; }
.stButton > button:hover, .stDownloadButton > button:hover { border-color: var(--primary); color: var(--primary); }
.stButton > button[kind="primary"], button[data-testid="stBaseButton-primary"] { background: var(--primary) !important; border-color: var(--primary) !important; color: var(--primary-fg) !important; }
.stButton > button[kind="primary"]:hover, button[data-testid="stBaseButton-primary"]:hover { background: #007a46 !important; border-color: #007a46 !important; color: var(--primary-fg) !important; }

.fi-call { display: flex; gap: 0.75rem; padding: 1rem 1.25rem; border-radius: var(--radius); border: 1px solid; margin: 1rem 0; font-size: 0.9rem; line-height: 1.6; }
.fi-call .ic { width: 24px; height: 24px; border-radius: 50%; flex: none; display: flex; align-items: center; justify-content: center; font-size: 0.75rem; font-weight: 700; margin-top: 2px; }
.fi-call b { font-weight: 600; }
.fi-call.info { background: rgba(59,130,246,0.1); border-color: rgba(59,130,246,0.3); } .fi-call.info .ic { background: rgba(59,130,246,0.2); color: #3b82f6; }
.fi-call.ok { background: rgba(0,138,80,0.1); border-color: rgba(0,138,80,0.3); } .fi-call.ok .ic { background: rgba(0,138,80,0.2); color: var(--primary); }
.fi-call.warn { background: rgba(245,158,11,0.1); border-color: rgba(245,158,11,0.3); } .fi-call.warn .ic { background: rgba(245,158,11,0.2); color: #f59e0b; }
.fi-call.err { background: rgba(239,68,68,0.1); border-color: rgba(239,68,68,0.3); } .fi-call.err .ic { background: rgba(239,68,68,0.2); color: #ef4444; }

.fi-err { background: rgba(239,68,68,0.1); border: 1px solid rgba(239,68,68,0.35); border-radius: var(--radius); padding: 1.5rem; margin: 1.5rem 0; }
.fi-err-h { display: flex; gap: 0.75rem; align-items: flex-start; }
.fi-err-h .ic { width: 32px; height: 32px; border-radius: 50%; background: rgba(239,68,68,0.2); color: #ef4444; font-weight: 700; display: flex; align-items: center; justify-content: center; flex: none; }
.fi-err .t { font-weight: 700; font-size: 1.125rem; }
.fi-err .w { color: var(--muted-fg); font-size: 0.875rem; margin-top: 0.25rem; }
.fi-err .w span { color: #ef4444; font-weight: 500; }
.fi-err-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 1rem; margin-top: 1rem; }
.fi-err-box { background: rgba(0,0,0,0.2); border: 1px solid var(--border); border-radius: var(--radius); padding: 1rem; }
.fi-err-box .k { font-size: 0.75rem; text-transform: uppercase; letter-spacing: 0.05em; color: var(--muted-fg); margin-bottom: 0.5rem; font-weight: 600; }
.fi-err-box .v { font-size: 0.9rem; line-height: 1.6; }
.fi-err-box.fix { border-color: rgba(0,138,80,0.3); } .fi-err-box.fix .k { color: var(--primary); }

.fi-checks { background: var(--card); border: 1px solid var(--border); border-radius: var(--radius); margin-bottom: 1.5rem; }
.fi-check { display: flex; gap: 0.75rem; padding: 1rem 1.25rem; border-top: 1px solid var(--border); }
.fi-check:first-child { border-top: none; }
.fi-check .ic { width: 24px; height: 24px; border-radius: 50%; flex: none; margin-top: 2px; display: flex; align-items: center; justify-content: center; font-size: 0.75rem; font-weight: 700; }
.fi-check.ok .ic { background: rgba(0,138,80,0.15); color: var(--primary); }
.fi-check.warn .ic { background: rgba(245,158,11,0.15); color: #f59e0b; }
.fi-check.err .ic { background: rgba(239,68,68,0.15); color: #ef4444; }
.fi-check .t { font-weight: 600; font-size: 0.95rem; }
.fi-check .d { color: var(--muted-fg); font-size: 0.875rem; margin-top: 0.25rem; word-break: break-word; }

.fi-tw { border: 1px solid var(--border); border-radius: var(--radius); overflow: auto; max-height: 450px; background: var(--card); margin: 1rem 0; }
.fi-table { width: 100%; border-collapse: collapse; font-size: 0.875rem; }
.fi-table th { position: sticky; top: 0; text-align: left; color: var(--muted-fg); font-weight: 600; font-size: 0.75rem; text-transform: uppercase; letter-spacing: 0.05em; padding: 0.75rem 1rem; border-bottom: 1px solid var(--border); background: var(--accent); z-index: 10; }
.fi-table td { padding: 0.75rem 1rem; border-bottom: 1px solid var(--border); vertical-align: middle; line-height: 1.5; }
.fi-table tr:last-child td { border-bottom: none; }
.fi-table td.row { font-family: monospace; color: #ef4444; white-space: nowrap; }
.fi-table td.val { font-family: monospace; color: #fca5a5; }
.fi-table td.fix { color: #6ee7b7; }
.fi-table td.mono { font-family: monospace; font-size: 0.8rem; color: var(--muted-fg); }
.fi-table td.status-ok { color: var(--primary); font-weight: 600; }
.fi-table td.status-fail { color: #ef4444; font-weight: 600; }

.fi-stats { display: grid; grid-template-columns: repeat(4, 1fr); gap: 1rem; margin: 1rem 0; }
.fi-stat { background: var(--card); border: 1px solid var(--border); border-radius: var(--radius); padding: 1rem; }
.fi-stat-l { color: var(--muted-fg); font-size: 0.75rem; text-transform: uppercase; letter-spacing: 0.05em; font-weight: 600; }
.fi-stat-v { font-size: 1.5rem; font-weight: 700; margin-top: 0.5rem; font-family: monospace; }
.fi-stat-v.sm { font-size: 1.125rem; }
.fi-stat-s { color: var(--muted-fg); font-size: 0.75rem; margin-top: 0.25rem; min-height: 1em; }
.fi-stat-v.g { color: var(--primary); }
.fi-stat-v.r { color: #ef4444; }
.fi-stat-v.a { color: #f59e0b; }
.fi-bar { height: 8px; background: var(--border); border-radius: 99px; overflow: hidden; margin-bottom: 1.5rem; display: flex; }
.fi-bar > .ok { height: 100%; background: var(--primary); transition: width 0.4s ease; }
.fi-bar > .bad { height: 100%; background: #ef4444; transition: width 0.4s ease; }

.fi-prev { display: grid; grid-template-columns: 1fr auto 1fr; gap: 1rem; align-items: center; background: var(--accent); border: 1px solid var(--border); border-radius: var(--radius); padding: 1rem 1.25rem; margin-top: 1rem; }
.fi-prev .k { font-size: 0.75rem; color: var(--muted-fg); text-transform: uppercase; letter-spacing: 0.05em; font-weight: 600; }
.fi-prev .v { font-family: monospace; font-size: 1rem; margin-top: 0.25rem; }
.fi-prev .arrow { color: var(--muted-fg); font-size: 1.25rem; }

.fi-calc { display: grid; grid-template-columns: 1fr 1fr; gap: 1rem; background: var(--accent); border: 1px solid var(--border); border-radius: var(--radius); padding: 1rem 1.25rem; margin-top: 1rem; }
.fi-calc .k { font-size: 0.75rem; color: var(--muted-fg); text-transform: uppercase; letter-spacing: 0.05em; font-weight: 600; }
.fi-calc .v { font-family: monospace; font-size: 1.25rem; font-weight: 700; margin-top: 0.25rem; color: var(--primary); }
.fi-calc .sub { font-size: 0.75rem; color: var(--muted-fg); margin-top: 0.25rem; }

@media (max-width: 768px) {
  .fi-stats, .fi-calc { grid-template-columns: repeat(2, 1fr); }
  .fi-err-grid { grid-template-columns: 1fr; }
  .fi-st .l { display: none; }
  .fi-st.active .l { display: inline; }
  .fi-tag { margin-left: 0; }
}
"""

STAGES = [("configure", "Setup", "Data & Form"), ("preflight", "Review", "Check Details"),
          ("execute", "Running", "Live Progress"), ("results", "Done", "Summary")]

STATUS_LABEL = {
    "waiting": ("Waiting", "a", "for start time"),
    "running": ("Running", "g", "in progress"),
    "done": ("Complete", "g", "all rows processed"),
    "stopped": ("Stopped", "a", "by user"),
    "error": ("Error", "r", "see details"),
}

def esc(v): return html.escape(str(v))
def md(html_text): st.markdown(html_text, unsafe_allow_html=True)
def inject_css(css): md(f"<style>{css}</style>")

def stat(label, value, sub="", tone="", small=False, raw_value=False):
    cls = f"fi-stat-v {tone} {'sm' if small else ''}"
    v = value if raw_value else esc(value)
    return (f'<div class="fi-stat"><div class="fi-stat-l">{esc(label)}</div>'
            f'<div class="{cls}">{v}</div><div class="fi-stat-s">{esc(sub)}</div></div>')

def stats_grid(*cards): md('<div class="fi-stats">' + "".join(cards) + "</div>")

def callout(kind, text_html):
    icon = {"info": "i", "ok": "✓", "warn": "!", "err": "✕"}[kind]
    md(f'<div class="fi-call {kind}"><span class="ic">{icon}</span><div>{text_html}</div></div>')

def html_table(columns, rows, classes=None):
    classes = classes or [""] * len(columns)
    head = "".join(f"<th>{esc(c)}</th>" for c in columns)
    body = "".join("<tr>" + "".join(f'<td class="{classes[i]}">{esc(v)}</td>' for i, v in enumerate(r)) + "</tr>" for r in rows)
    md(f'<div class="fi-tw"><table class="fi-table"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')

def error_card(title, where="", why="", fix="", problems=None, technical=""):
    where_html = f'<div class="w">Where: <span>{esc(where)}</span></div>' if where else ""
    boxes = ""
    if why or fix:
        boxes = ('<div class="fi-err-grid">'
                 f'<div class="fi-err-box"><div class="k">Why it happened</div><div class="v">{esc(why)}</div></div>'
                 f'<div class="fi-err-box fix"><div class="k">How to fix it</div><div class="v">{esc(fix)}</div></div>'
                 '</div>')
    md(f'<div class="fi-err"><div class="fi-err-h"><span class="ic">!</span>'
       f'<div><div class="t">{esc(title)}</div>{where_html}</div></div>{boxes}</div>')
    if problems:
        html_table(["Row", "Column", "Value", "Problem", "Fix"],
                   [[p.row, p.column, p.value, p.problem, p.fix] for p in problems[:300]],
                   ["row", "", "val", "", "fix"])
        if len(problems) > 300: st.caption(f"Showing first 300 of {len(problems)} problems.")
    if technical:
        with st.expander("Technical Details"): st.code(technical, language=None)

def show_error(exc): error_card(exc.title, exc.where, exc.why, exc.fix, exc.problems, exc.technical)

def stepper(stage):
    idx = [k for k, _, _ in STAGES].index(stage)
    parts = []
    for i, (_, label, sub) in enumerate(STAGES):
        cls = "done" if i < idx else "active" if i == idx else ""
        mark = "✓" if i < idx else str(i + 1)
        parts.append(f'<div class="fi-st {cls}"><span class="n">{mark}</span>'
                     f'<span class="l">{label}<small>{sub}</small></span></div>')
        if i < len(STAGES) - 1: parts.append(f'<div class="fi-st-line {"done" if i < idx else ""}"></div>')
    md('<div class="fi-stepper">' + "".join(parts) + "</div>")

def section_head(title, desc): md(f'<div class="fi-sec"><h2>{esc(title)}</h2><p>{esc(desc)}</p></div>')
def step_head(num, title, desc): md(f'<div class="fi-step"><b>{num}</b><span class="t">{esc(title)}</span></div><div class="fi-step-d">{esc(desc)}</div>')

def format_mm_ss(seconds):
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"

def format_human(seconds):
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    if h: return f"{h}h {m}m"
    if m: return f"{m}m {s}s"
    return f"{s}s"

def fmt_clock(dt): 
    return dt.strftime("%I:%M:%S %p").lstrip("0")

# Custom 12-hour AM/PM time picker formatted as HH:MM:SS AM/PM
def get_time_12h(label, default_time, key_base):
    with st.container(border=True):
        st.markdown(f"**{label}**")
        
        if default_time is None:
            default_time = datetime.now(IST).time()
            
        h, m, s = default_time.hour, default_time.minute, default_time.second
        period = "AM" if h < 12 else "PM"
        h12 = h % 12
        if h12 == 0: h12 = 12
        
        cols = st.columns([2, 1, 2, 1, 2, 1, 3])
        with cols[0]:
            hour = st.number_input("Hour", 1, 12, h12, key=f"{key_base}_h", label_visibility="collapsed")
        with cols[1]:
            st.markdown("<div style='padding-top:1.5rem;text-align:center;font-weight:bold'>:</div>", unsafe_allow_html=True)
        with cols[2]:
            minute = st.number_input("Min", 0, 59, m, key=f"{key_base}_m", label_visibility="collapsed")
        with cols[3]:
            st.markdown("<div style='padding-top:1.5rem;text-align:center;font-weight:bold'>:</div>", unsafe_allow_html=True)
        with cols[4]:
            second = st.number_input("Sec", 0, 59, s, key=f"{key_base}_s", label_visibility="collapsed")
        with cols[5]:
            st.markdown("<div style='padding-top:1.5rem;text-align:center;font-weight:bold'></div>", unsafe_allow_html=True)
        with cols[6]:
            ampm = st.selectbox("AM/PM", ["AM", "PM"], index=0 if period=="AM" else 1, key=f"{key_base}_p", label_visibility="collapsed")
            
        final_h = hour
        if ampm == "PM" and hour != 12: final_h += 12
        if ampm == "AM" and hour == 12: final_h = 0
        
        return dtime(final_h, minute, second)

# ============================================================
#  ERRORS & DATA LOADING
# ============================================================

@dataclass
class Problem:
    row: int | None; column: str; value: str; problem: str; fix: str

class AppError(Exception):
    def __init__(self, title, *, where="", why="", fix="", problems=None, technical=""):
        super().__init__(title)
        self.title, self.where, self.why, self.fix = title, where, why, fix
        self.problems = problems or []
        self.technical = technical

def explain_exception(exc):
    tech = "".join(traceback.format_exception(type(exc), exc, exc.__traceback__))[-3000:]
    if isinstance(exc, PermissionError): return dict(title="Cannot save file", why="File is open or access denied.", fix="Close the file and try again.", technical=tech)
    if isinstance(exc, OSError) and getattr(exc, "errno", None) == 28: return dict(title="Out of disk space", why="No room to save.", fix="Try again later.", technical=tech)
    return dict(title="Unexpected error", why="Unknown error occurred.", fix="Contact support.", technical=tech)

def sheet_row(idx): return int(idx) + 2

def sheet_csv_url(url):
    url = url.strip()
    if re.match(r"https://docs\.google\.com/spreadsheets/d/e/", url):
        if "output=csv" in url: return url
        base = url.split("?")[0].split("#")[0]
        if base.endswith("/pubhtml"): base = base[: -len("/pubhtml")] + "/pub"
        return base + "?output=csv"
    m = re.match(r"https://docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)", url)
    if not m: return None
    out = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    gid = re.search(r"[#&?]gid=(\d+)", url)
    if gid: out += f"&gid={gid.group(1)}"
    return out

def _read_file(path, name):
    lower = name.lower()
    if lower.endswith((".xlsx", ".xls")): return pd.read_excel(path)
    if lower.endswith(".ods"): return pd.read_excel(path, engine="odf")
    if lower.endswith(".tsv"): return pd.read_csv(path, sep="\t")
    if lower.endswith(".csv"): return pd.read_csv(path)
    if lower.endswith(".json"): return pd.read_json(path)
    raise AppError("Unsupported file type", where="Step 1", why=f"'{name}' is not supported.", fix="Use .xlsx or .csv.")

def load_dataframe(inp, workdir, log):
    kind = inp["source_kind"]
    if kind == "sheet":
        raw = inp["sheet_url"].strip()
        if not raw: raise AppError("No Google Sheet link", where="Step 1", why="Box is empty.", fix="Paste link.")
        url = sheet_csv_url(raw)
        if not url: raise AppError("Invalid Sheets link", where="Step 1", why="Not a valid Sheets URL.", fix="Copy from address bar.", technical=raw)
        try: resp = requests.get(url, timeout=30)
        except requests.exceptions.RequestException as exc: raise AppError("Connection failed", where="Step 1", why="Could not reach Google.", fix="Check internet.", technical=str(exc))
        if resp.status_code != 200 or resp.text.lstrip().startswith("<"): raise AppError("Sheet not public", where="Step 1", why="Google asked for login.", fix="Set sharing to 'Anyone with link'.", technical=f"HTTP {resp.status_code}")
        df = pd.read_csv(io.StringIO(resp.text))
        log(f"Google Sheet loaded: {len(df)} rows")
        return df, "Google Sheet", os.path.join(workdir, "responses_updated.xlsx")
    if kind == "manual":
        text = inp.get("manual_data", "").strip()
        if not text: raise AppError("No data entered", where="Step 1", why="Text box empty.", fix="Paste CSV data.")
        try: df = pd.read_csv(io.StringIO(text))
        except Exception as exc: raise AppError("Invalid CSV", where="Step 1", why="Could not parse text.", fix="Check format.", technical=str(exc))
        if df.empty: raise AppError("No rows", where="Step 1", why="Only headers found.", fix="Add data rows.")
        log(f"Manual data loaded: {len(df)} rows")
        return df, "Manual Entry", os.path.join(workdir, "manual_responses_updated.xlsx")
    if kind == "upload":
        if not inp.get("upload"): raise AppError("No file", where="Step 1", why="Nothing uploaded.", fix="Upload file.")
        name, data = inp["upload"]
        name = os.path.basename(name)
        path = os.path.join(workdir, name)
        with open(path, "wb") as fh: fh.write(data)
    else:
        path = inp["local_path"].strip().strip('"').strip("'")
        name = os.path.basename(path)
        if not path: raise AppError("No path", where="Step 1", why="Path empty.", fix="Enter path.")
        if not os.path.isfile(path): raise AppError("File not found", where="Step 1", why=f"No file at '{path}'.", fix="Check path.")
    try: df = _read_file(path, name)
    except AppError: raise
    except Exception as exc: raise AppError("Cannot open file", where=f"Step 1 · {name}", why="File may be damaged.", fix="Save as .xlsx/.csv.", technical=str(exc))
    stem, ext = os.path.splitext(name)
    if ext.lower() == ".xls": ext = ".xlsx"
    out_dir = workdir if kind == "upload" else os.path.dirname(os.path.abspath(path))
    return df, name, os.path.join(out_dir, f"{stem}_updated{ext}")

def save_dataframe(df, file_path):
    root, ext = os.path.splitext(file_path)
    tmp = f"{root}.saving{ext}"
    lower = ext.lower()
    if lower == ".xlsx": df.to_excel(tmp, index=False)
    elif lower == ".tsv": df.to_csv(tmp, sep="\t", index=False)
    elif lower == ".json": df.to_json(tmp, orient="records", indent=2)
    elif lower == ".ods": df.to_excel(tmp, index=False, engine="odf")
    else: df.to_csv(tmp, index=False)
    os.replace(tmp, file_path)

# ============================================================
#  FORM METADATA & VALIDATION
# ============================================================

def extract_entry_ids(prefilled_link):
    parsed = urllib.parse.urlparse(prefilled_link)
    pairs = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    seen, ids = set(), []
    for key, _ in pairs:
        if key.startswith("entry.") and key not in seen: ids.append(key); seen.add(key)
    return ids

def _parse_fb_blob(html_text):
    token = "FB_PUBLIC_LOAD_DATA_"
    idx = html_text.find(token)
    if idx == -1: return None
    start = html_text.find("[", idx)
    if start == -1: return None
    depth, in_str, escape = 0, False, False
    for i in range(start, len(html_text)):
        ch = html_text[i]
        if in_str:
            if escape: escape = False
            elif ch == "\\": escape = True
            elif ch == '"': in_str = False
            continue
        if ch == '"': in_str = True
        elif ch == "[": depth += 1
        elif ch == "]":
            depth -= 1
            if depth == 0:
                try: return json.loads(html_text[start:i + 1])
                except json.JSONDecodeError: return None
    return None

@dataclass
class FormQuestion:
    entry: str; label: str; options: list; required: bool; grid_row: str = ""

def _questions_from_blob(blob):
    out = []
    try: items = blob[1][1]
    except (IndexError, TypeError): return out
    for item in items or []:
        if not isinstance(item, list) or len(item) < 5: continue
        label = item[1] if isinstance(item[1], str) else ""
        question = item[4]
        if not isinstance(question, list): continue
        is_grid = sum(1 for q in question if isinstance(q, list) and q and q[0] is not None) >= 2
        for q_item in question:
            if not isinstance(q_item, list) or not q_item or q_item[0] is None: continue
            opts = []
            if len(q_item) > 1 and isinstance(q_item[1], list):
                opts = [o[0] for o in q_item[1] if isinstance(o, list) and o and isinstance(o[0], str) and o[0]]
            required = len(q_item) > 2 and q_item[2] in (1, True)
            sub = ""
            if is_grid:
                lbl = q_item[3] if len(q_item) > 3 else None
                if isinstance(lbl, list) and lbl and isinstance(lbl[0], str): lbl = lbl[0]
                sub = lbl if isinstance(lbl, str) and lbl else f"row {q_item[0]}"
            out.append(FormQuestion(f"entry.{q_item[0]}", label, opts, required, sub))
    return out

@dataclass
class FormInfo:
    field_types: dict = field(default_factory=dict); options: dict = field(default_factory=dict)
    labels: dict = field(default_factory=dict); required: set = field(default_factory=set); reachable: bool = False

def fetch_form_metadata(prefilled_link):
    info = FormInfo()
    try: resp = requests.get(prefilled_link, timeout=30)
    except requests.exceptions.RequestException: return info
    if resp.status_code != 200: return info
    info.reachable = True
    page = resp.text
    for m in re.finditer(r'name="(entry\.\d+)"[^>]*type="([^"]+)"', page):
        eid, itype = m.group(1), m.group(2).lower()
        if eid not in info.field_types: info.field_types[eid] = itype
        elif "checkbox" in (info.field_types[eid], itype): info.field_types[eid] = "checkbox"
    for m in re.finditer(r'<textarea[^>]*name="(entry\.\d+)"', page): info.field_types.setdefault(m.group(1), "textarea")
    for m in re.finditer(r'<select[^>]*name="(entry\.\d+)"', page): info.field_types.setdefault(m.group(1), "select")
    blob = _parse_fb_blob(page)
    if blob:
        questions = _questions_from_blob(blob)
        for q in questions:
            if q.options: info.options[q.entry] = q.options
            info.labels[q.entry] = f"{q.label} — {q.grid_row}" if q.grid_row else q.label
            if q.required: info.required.add(q.entry)
            if q.grid_row: info.field_types[q.entry] = "grid"
    return info

def normalize_value(raw, entry_type="text"):
    if pd.isna(raw): return BLANK
    if isinstance(raw, (pd.Timestamp, datetime)):
        if entry_type == "date": return raw.strftime("%Y-%m-%d")
        if entry_type == "time": return raw.strftime("%H:%M")
        if entry_type == "datetime-local": return raw.strftime("%Y-%m-%dT%H:%M")
    if isinstance(raw, float) and raw.is_integer(): raw = int(raw)
    text = str(raw).strip()
    if not text: return BLANK
    if entry_type == "datetime-local" and " " in text and "T" not in text: text = text.replace(" ", "T", 1)
    return text

def is_blank(value): return value == BLANK

def best_option(value, options):
    if value in options: return value
    squash = lambda s: re.sub(r"\s+", " ", str(s)).strip().lower()
    by_key = {squash(o): o for o in options}
    if squash(value) in by_key: return by_key[squash(value)]
    try:
        num = float(value)
        for o in options:
            try:
                if float(o) == num: return o
            except (TypeError, ValueError): pass
    except ValueError: pass
    close = difflib.get_close_matches(squash(value), list(by_key), n=1, cutoff=0.85)
    return by_key[close[0]] if close else None

def split_choices(value, options):
    if value in options: return [value]
    return [v.strip() for v in value.split(",") if v.strip()]

def _opts_hint(opts, limit=8):
    shown = ", ".join(f"'{o}'" for o in opts[:limit])
    return shown + (f" (+{len(opts) - limit} more)" if len(opts) > limit else "")

def auto_correct(df, pending_indices, required_columns, form: FormInfo):
    corrections, problems = [], []
    for idx in pending_indices:
        for col, entry in required_columns.items():
            etype = form.field_types.get(entry, "text")
            value = normalize_value(df.at[idx, col], etype)
            opts = form.options.get(entry)
            new_val, reason = value, None
            if is_blank(value):
                if entry in form.required: problems.append(Problem(sheet_row(idx), col, "(empty)", "Required question is empty.", f"Fill in an answer{': ' + _opts_hint(opts) if opts else ''}."))
                continue
            if opts:
                if etype == "checkbox":
                    parts = split_choices(value, opts)
                    fixed = [best_option(p, opts) for p in parts]
                    bad = [p for p, f in zip(parts, fixed) if f is None]
                    if bad: problems.append(Problem(sheet_row(idx), col, value, f"{', '.join(f'{b}' for b in bad)} not in choices.", f"Use: {_opts_hint(opts)}.")); continue
                    candidate = ", ".join(dict.fromkeys(fixed))
                    if candidate != value: new_val, reason = candidate, "matched choices"
                else:
                    candidate = best_option(value, opts)
                    if candidate is None: problems.append(Problem(sheet_row(idx), col, value, "Answer not in choices.", f"Use one of: {_opts_hint(opts)}.")); continue
                    if candidate != value: new_val, reason = candidate, "matched choice"
            if etype in ("date", "time", "datetime-local"):
                fmt, what, example = {"date": ("%Y-%m-%d", "date", "2026-10-29"), "time": ("%H:%M", "time", "14:30"), "datetime-local": ("%Y-%m-%dT%H:%M", "date/time", "2026-10-29 14:30")}[etype]
                try:
                    new_val = pd.to_datetime(value).strftime(fmt)
                    if new_val != value: reason = f"{what} rewritten"
                except (ValueError, TypeError, OverflowError): problems.append(Problem(sheet_row(idx), col, value, f"Invalid {what}.", f"Write like {example}.")); continue
            if new_val != value:
                df.at[idx, col] = new_val
                corrections.append(dict(Row=sheet_row(idx), Column=col, From=value, To=new_val, Why=reason))
    return corrections, problems

def validate_data(df, pending_indices, required_columns, form: FormInfo):
    problems = []
    for idx in pending_indices:
        for col, entry in required_columns.items():
            etype = form.field_types.get(entry, "text")
            value = normalize_value(df.at[idx, col], etype)
            if is_blank(value): continue
            opts = form.options.get(entry)
            if opts:
                parts = split_choices(value, opts) if etype == "checkbox" else [value]
                bad = [p for p in parts if p not in opts]
                if bad: problems.append(Problem(sheet_row(idx), col, value, "Answer not in choices.", f"Use: {_opts_hint(opts)}."))
    return problems

# ============================================================
#  SCHEDULING & SUBMISSION
# ============================================================

def base_window(start_t, end_t):
    today = datetime.now(IST).date()
    start_dt = IST.localize(datetime.combine(today, start_t))
    end_dt = IST.localize(datetime.combine(today, end_t))
    if end_dt <= start_dt: end_dt += timedelta(days=1)
    return start_dt, end_dt

def compute_time_window(start_t, end_t, start_range, end_range, total_rows, log):
    min_required_sec = total_rows * MIN_SECONDS_PER_ROW
    start_dt, end_dt = base_window(start_t, end_t)
    original_window_sec = int((end_dt - start_dt).total_seconds())
    srmin, srmax = sorted(start_range)
    ermin, ermax = sorted(end_range)
    rs = re_ = window_sec = 0
    rand_start = rand_end = None
    for _ in range(MAX_SCHEDULE_ATTEMPTS):
        rs = random.randint(srmin, srmax)
        re_ = random.randint(ermin, ermax)
        rand_start = start_dt + timedelta(seconds=rs)
        rand_end = end_dt + timedelta(seconds=re_)
        if rand_end <= rand_start: rand_end = rand_start + timedelta(seconds=1)
        window_sec = int((rand_end - rand_start).total_seconds())
        if window_sec >= original_window_sec and window_sec >= min_required_sec: break
    else:
        if original_window_sec < min_required_sec: raise AppError("Window too short", where="Step 4", why=f"{total_rows} rows need {format_human(min_required_sec)}.", fix="Increase window.")
        raise AppError("Random delays invalid", where="Step 4", why="Offsets make window shorter than entered.", fix="Adjust ranges.")
    log("========== TIME CONFIGURATION ==========")
    log(f"Entered start: {start_dt.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Entered end: {end_dt.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Actual start: {rand_start.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Actual end: {rand_end.strftime('%d-%m-%Y %I:%M:%S %p IST')}")
    log(f"Rows: {total_rows}")
    log("========================================")
    cfg = dict(entered_start=start_dt, entered_end=end_dt, rs=rs, re=re_, rand_start=rand_start, rand_end=rand_end, original_window=original_window_sec, window=window_sec, min_required=min_required_sec, total_rows=total_rows)
    return rand_start, rand_end, window_sec, cfg

def wait_until(target_dt, stop):
    while True:
        remaining = (target_dt - datetime.now(IST)).total_seconds()
        if remaining <= 0: return True
        if stop.wait(min(remaining, 1.0)): return False

@dataclass
class SubmitResult:
    ok: bool; http: int; problem: str = ""; fix: str = ""; technical: str = ""

def build_payload(row, required_columns, form: FormInfo):
    payload = {}
    for col, entry in required_columns.items():
        etype = form.field_types.get(entry, "text")
        value = normalize_value(row[col], etype)
        opts = form.options.get(entry)
        if is_blank(value) and opts: continue
        if etype == "checkbox": payload[entry] = split_choices(value, opts or [])
        else: payload[entry] = value
    return payload

def explain_response(status, text, url):
    snippet = text[:300] if text else ""
    tech = f"HTTP {status} · {url}\n{snippet}"
    low = (text or "").lower()
    if status == 200:
        if "accounts.google.com" in url or "servicelogin" in low: return SubmitResult(False, status, "Form requires sign-in.", "Turn off 'Restrict to users'.", tech)
        if "no longer accepting responses" in low: return SubmitResult(False, status, "Form closed.", "Enable 'Accepting responses'.", tech)
        return SubmitResult(True, status)
    if status == 0: return SubmitResult(False, status, "Connection failed.", "Check internet.", tech)
    if status == 400: return SubmitResult(False, status, "Google rejected answers.", "Check required fields/choices.", tech)
    if status in (401, 403): return SubmitResult(False, status, "Access denied.", "Check form permissions.", tech)
    if status == 404: return SubmitResult(False, status, "Form not found.", "Check link.", tech)
    if status == 429: return SubmitResult(False, status, "Too many requests.", "Increase time window.", tech)
    if status >= 500: return SubmitResult(False, status, "Google server error.", "Try again later.", tech)
    return SubmitResult(False, status, f"Unexpected error ({status}).", "Try again.", tech)

def submit_row(session, row, prep, stop, max_retries=3):
    payload = build_payload(row, prep.required_columns, prep.form)
    headers = {"User-Agent": "Mozilla/5.0", "Referer": prep.prefilled_link, "Origin": "https://docs.google.com"}
    last = SubmitResult(False, 0, technical="no attempt")
    for attempt in range(max_retries):
        try: r = session.post(prep.form_url, data=payload, headers=headers, timeout=30)
        except requests.exceptions.RequestException as exc:
            last = explain_response(0, "", prep.form_url)
            last.technical = f"{type(exc).__name__}: {exc}"
            delay = 2 * (2 ** attempt)
        else:
            if r.status_code != 429 and r.status_code < 500: return explain_response(r.status_code, r.text, r.url)
            last = explain_response(r.status_code, r.text, r.url)
            try: delay = float(r.headers.get("Retry-After", ""))
            except ValueError: delay = 2 * (2 ** attempt)
        if attempt < max_retries - 1 and stop.wait(min(delay, 60)): break
    last.technical = f"Tried {max_retries} times.\n{last.technical}"
    return last

def generate_analysis_sheet(df, file_path, submitted_count, log):
    if not OPENPYXL_AVAILABLE: log("Analysis skipped: openpyxl missing."); return
    if not file_path.lower().endswith(".xlsx"): log("Analysis skipped: not .xlsx."); return
    wb = load_workbook(file_path)
    if "Analysis" in wb.sheetnames: del wb["Analysis"]
    ws = wb.create_sheet("Analysis")
    NAVY, BLUE, LIGHT_BLUE, WHITE = "1F3864", "2E75B6", "DEEAF1", "FFFFFF"
    def solid(color): return PatternFill("solid", fgColor=color)
    thin = Side(style="thin", color="CCCCCC")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:F1")
    ws["A1"] = "Response Analysis"
    ws["A1"].fill = solid(NAVY); ws["A1"].font = Font(bold=True, color=WHITE, size=14); ws["A1"].alignment = ctr
    ws.row_dimensions[1].height = 28
    ws["A2"] = f"Generated: {datetime.now(IST).strftime('%d-%m-%Y %I:%M:%S %p')}"
    ws["B2"] = f"Total rows: {len(df)}"
    ws["C2"] = f"Submitted: {submitted_count}"
    analysis_cols = [c for c in df.columns if str(c).strip().lower() not in STATUS_NAMES]
    current_row = 4
    for col_name in analysis_cols:
        series = df[col_name].dropna().astype(str).str.strip()
        series = series[series != ""]
        if series.empty: continue
        is_checkbox = series.str.contains(",").any()
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        cell = ws.cell(row=current_row, column=1, value=str(col_name))
        cell.fill = solid(BLUE); cell.font = Font(bold=True, color=WHITE, size=11); cell.alignment = ctr
        current_row += 1
        hdr_row = current_row
        for ci, h in enumerate(["#", "Answer", "Count", "% of Total"], 1):
            c = ws.cell(row=current_row, column=ci, value=h)
            c.fill = solid(BLUE); c.font = Font(bold=True, color=WHITE, size=10); c.alignment = ctr; c.border = bdr
        current_row += 1
        if is_checkbox:
            all_vals = []
            for v in series: all_vals.extend(x.strip() for x in v.split(",") if x.strip())
            freq = Counter(all_vals)
        else: freq = Counter(series.tolist())
        total_responders = len(series)
        items_sorted = sorted(freq.items(), key=lambda x: -x[1])
        data_start = current_row
        for ri, (answer, count) in enumerate(items_sorted, 1):
            pct = (count / total_responders * 100) if total_responders else 0
            for ci, val in enumerate([ri, answer, count, f"{pct:.1f}%"], 1):
                c = ws.cell(row=current_row, column=ci, value=val)
                c.fill = solid(LIGHT_BLUE if ri % 2 == 0 else WHITE); c.font = Font(color="1F1F1F", size=10); c.alignment = ctr; c.border = bdr
            current_row += 1
        total_row_idx = current_row
        total_pct = "100%" if not is_checkbox else "---"
        for ci, val in enumerate(["", "TOTAL", sum(freq.values()), total_pct], 1):
            c = ws.cell(row=current_row, column=ci, value=val)
            c.fill = solid(NAVY); c.font = Font(bold=True, color=WHITE, size=10); c.alignment = ctr; c.border = bdr
        current_row += 1
        if items_sorted:
            chart = BarChart(); chart.type = "col"; chart.grouping = "clustered"
            chart.title = str(col_name)[:30]; chart.y_axis.title = "Count"; chart.x_axis.title = "Answer"
            chart.style = 10; chart.width = 18; chart.height = 12
            data_ref = Reference(ws, min_col=3, max_col=3, min_row=hdr_row, max_row=total_row_idx - 1)
            cats_ref = Reference(ws, min_col=2, max_col=2, min_row=data_start, max_row=total_row_idx - 1)
            chart.add_data(data_ref, titles_from_data=True); chart.set_categories(cats_ref)
            ws.add_chart(chart, f"{get_column_letter(7)}{data_start}")
        current_row += 2
    wb.save(file_path)
    log("Analysis sheet added.")

# ============================================================
#  PREPARATION & JOB
# ============================================================

@dataclass
class Prepared:
    df: pd.DataFrame; output_path: str; status_col: str; required_columns: dict; form: FormInfo
    pending_indices: list; form_url: str; prefilled_link: str; rand_start: datetime; rand_end: datetime
    window_sec: int; pre_lines: list; schedule_lines: list; inp: dict; summary: dict; mapping: list; corrections: list; time_cfg: dict

def check_link(link):
    link = link.strip()
    how = "Open form → ⋮ menu → 'Get pre-filled link' → Fill answers → 'Get link' → Copy."
    if not link: raise AppError("No form link", where="Step 2", why="Box empty.", fix=how)
    if "/viewform" not in link or not link.startswith("http"): raise AppError("Invalid link", where="Step 2", why="Not a pre-filled link.", fix=how, technical=link)
    if not extract_entry_ids(link): raise AppError("No entry IDs", where="Step 2", why="Link missing 'entry.' values.", fix=how, technical=link)
    return link

def map_columns(df, prefilled_link, status_col):
    entry_cols = [c for c in df.columns if str(c).startswith("entry.")]
    prefilled_ids = extract_entry_ids(prefilled_link)
    non_status_cols = [c for c in df.columns if c != status_col]
    if entry_cols: return {c: c for c in entry_cols}
    if prefilled_ids:
        matched = [c for c in prefilled_ids if c in df.columns]
        if matched: return {c: c for c in matched}
        if len(prefilled_ids) == len(non_status_cols): return dict(zip(non_status_cols, prefilled_ids))
    raise AppError("Columns mismatch", where="Step 1 & 2", why=f"Sheet has {len(non_status_cols)} cols, form has {len(prefilled_ids)} questions.", fix="Match counts or rename columns to entry.IDs.", technical=f"Cols: {list(map(str, non_status_cols))}\nIDs: {prefilled_ids}")

def prepare(inp: dict, workdir: str) -> Prepared:
    lines: list = []
    log = lines.append
    prefilled_link = check_link(inp["link"])
    form_url = prefilled_link.split("/viewform")[0] + "/formResponse"
    df, source_label, output_path = load_dataframe(inp, workdir, log)
    if df.empty: raise AppError("No rows", where="Step 1", why="Only headers found.", fix="Add data.")
    df.columns = [str(c).strip() for c in df.columns]
    n_rows, n_cols = len(df), len(df.columns)
    status_col = next((c for c in df.columns if c.lower() in STATUS_NAMES), None)
    if status_col is None: status_col = "Status"; df[status_col] = ""
    required_columns = map_columns(df, prefilled_link, status_col)
    status_series = df[status_col].fillna("").astype(str).str.strip().str.lower()
    pending_indices = df[status_series != "submitted"].index.tolist()
    pending_total = len(pending_indices)
    if pending_total == 0: raise AppError("All submitted", where="Step 1", why="No pending rows.", fix=f"Clear '{status_col}' column.")
    if inp["random"]:
        random.shuffle(pending_indices)
        n = int(inp["sample_n"] or 0)
        if n > 0: pending_indices = pending_indices[: min(n, len(pending_indices))]
    total_rows = len(pending_indices)
    log(f"Form link: {prefilled_link}")
    log("Column → question:")
    for col, entry in required_columns.items(): log(f"  {col} → {entry}")
    form = fetch_form_metadata(prefilled_link)
    if form.field_types:
        log("Question types read:")
        for entry in required_columns.values(): log(f"  {entry}: {form.field_types.get(entry, 'unknown')}")
    else: log("Could not read question types.")
    corrections, problems = auto_correct(df, pending_indices, required_columns, form)
    if corrections:
        log(f"Auto-fixed {len(corrections)} values:")
        for c in corrections: log(f"  Row {c['Row']} | {c['Column']} | {c['From']} → {c['To']} ({c['Why']})")
    if not problems: problems = validate_data(df, pending_indices, required_columns, form)
    if problems:
        rows = sorted({p.row for p in problems})
        raise AppError(f"{len(problems)} cells need fixing", where=f"Sheet rows {', '.join(map(str, rows[:10]))}", why="Answers don't match form.", fix="Fix cells and upload again.", problems=problems)
    if corrections: save_dataframe(df, output_path)
    sched_lines: list = []
    rand_start, rand_end, window_sec, time_cfg = compute_time_window(inp["start"], inp["end"], inp["srange"], inp["erange"], total_rows, sched_lines.append)
    summary = dict(source=source_label, is_google_sheets=inp["source_kind"] == "sheet", rows=n_rows, cols=n_cols, pending=pending_total, selected=total_rows, random=inp["random"], entry_fields=len(extract_entry_ids(prefilled_link)), validated_fields=len(required_columns), form_reachable=form.reachable)
    mapping = [dict(Column=col, Question=form.labels.get(entry) or "—", ID=entry, Type=form.field_types.get(entry, "unknown"), Required="Yes" if entry in form.required else "") for col, entry in required_columns.items()]
    return Prepared(df=df, output_path=output_path, status_col=status_col, required_columns=required_columns, form=form, pending_indices=pending_indices, form_url=form_url, prefilled_link=prefilled_link, rand_start=rand_start, rand_end=rand_end, window_sec=window_sec, pre_lines=lines, schedule_lines=sched_lines, inp=inp, summary=summary, mapping=mapping, corrections=corrections, time_cfg=time_cfg)

def reroll_schedule(prep: Prepared):
    lines: list = []
    i = prep.inp
    rs_, re_, ws, cfg = compute_time_window(i["start"], i["end"], i["srange"], i["erange"], prep.summary["selected"], lines.append)
    prep.rand_start, prep.rand_end, prep.window_sec = rs_, re_, ws
    prep.time_cfg, prep.schedule_lines = cfg, lines

def console_text(prep: Prepared) -> str:
    return "\n".join(prep.pre_lines + [""] + prep.schedule_lines).strip()

class Job:
    def __init__(self, prep: Prepared):
        self.prep = prep; self.total = len(prep.pending_indices); self.lock = threading.Lock()
        self.stop = threading.Event(); self.status = "waiting"; self.rows: list = []; self.failures: list = []
        self.submitted = 0; self.failed = 0; self.next_target = None; self.error = None; self.notified = False
        self.t_start = None; self.t_end = None; self.thread = threading.Thread(target=self._run, daemon=True)
    def start(self): self.thread.start()
    @property
    def active(self): return self.status in ("waiting", "running")
    def _set(self, **kw):
        with self.lock:
            for k, v in kw.items(): setattr(self, k, v)
    def _finish(self, status): self._set(t_end=time.time(), status=status, next_target=None)
    def elapsed(self):
        if not self.t_start: return 0
        return int((self.t_end or time.time()) - self.t_start)
    def _add(self, **row):
        with self.lock: self.rows.append(row)
    def _note(self, text): self._add(kind="note", text=text)
    def snapshot(self):
        with self.lock:
            return {"status": self.status, "rows": list(self.rows), "failures": list(self.failures), "submitted": self.submitted, "failed": self.failed, "total": self.total, "next_target": self.next_target, "error": self.error}
    def _run(self):
        p = self.prep; session = requests.Session(); dirty, last_save = False, time.time()
        try:
            if datetime.now(IST) < p.rand_start:
                self._set(next_target=p.rand_start)
                self._note(f"Waiting for start time ({fmt_clock(p.rand_start)} IST)")
                if not wait_until(p.rand_start, self.stop): self._note("Stopped before start."); self._finish("stopped"); return
            self._set(t_start=time.time(), status="running")
            total_rows = self.total; window_sec = max(p.window_sec, total_rows + 5)
            offsets = sorted(random.sample(range(1, window_sec + 1), total_rows))
            start_ref = datetime.now(IST)
            for i, (offset, idx) in enumerate(zip(offsets, p.pending_indices), 1):
                if self.stop.is_set(): break
                target = start_ref + timedelta(seconds=offset)
                self._set(next_target=target)
                if not wait_until(target, self.stop): break
                res = submit_row(session, p.df.loc[idx], p, self.stop)
                now_s = fmt_clock(datetime.now(IST)); row_no = sheet_row(idx)
                if res.ok:
                    p.df.at[idx, p.status_col] = "submitted"; dirty = True
                    with self.lock: self.submitted += 1
                    if i < total_rows:
                        nxt = start_ref + timedelta(seconds=offsets[i])
                        gap_sec = offsets[i] - offsets[i - 1]
                        self._add(kind="ok", row=row_no, time=now_s, nxt=fmt_clock(nxt), gap_sec=gap_sec, text=f"Row {row_no} sent")
                    else: self._add(kind="final", row=row_no, time=now_s, text=f"Row {row_no} sent (last)")
                else:
                    fail = dict(row=row_no, time=now_s, problem=res.problem, fix=res.fix, technical=res.technical, http=res.http)
                    with self.lock: self.failed += 1; self.failures.append(fail)
                    self._add(kind="fail", **fail, text=f"Row {row_no} failed")
                if dirty and (time.time() - last_save >= SAVE_EVERY_SEC or i == total_rows):
                    save_dataframe(p.df, p.output_path); dirty, last_save = False, time.time()
            if dirty: save_dataframe(p.df, p.output_path)
            if self.stop.is_set(): self._note("Stopped by user."); self._finish("stopped"); return
            self._note("All rows processed.")
            try: generate_analysis_sheet(p.df, p.output_path, self.submitted, self._note)
            except Exception as exc: self._note(f"Analysis sheet failed ({type(exc).__name__}).")
            self._finish("done")
        except Exception as exc:
            if dirty:
                try: save_dataframe(p.df, p.output_path)
                except Exception: pass
            err = explain_exception(exc); self._set(error=err); self._note(f"ERROR: {err['title']}"); self._finish("error")
        finally: session.close()
    def log_text(self):
        lines = list(self.prep.pre_lines) + [""] + list(self.prep.schedule_lines) + [""]
        with self.lock:
            lines += [r["text"] for r in self.rows]
            if self.failures:
                lines += ["", "========== FAILED ROWS =========="]
                for f in self.failures: lines += [f"Row {f['row']}: {f['problem']}", f"  Fix: {f['fix']}", f"  Tech: {f['technical']}"]
            if self.error: lines += ["", "========== ERROR ==========", self.error["title"], self.error["technical"]]
        return "\n".join(lines)
    def failures_csv(self):
        with self.lock:
            rows = [dict(**{"Row": f["row"], "Time": f["time"], "Problem": f["problem"], "Fix": f["fix"], "Technical": f["technical"]}) for f in self.failures]
        return pd.DataFrame(rows).to_csv(index=False)

# ============================================================
#  UI VIEWS
# ============================================================

def schedule_preview(start_t, end_t, srange, erange):
    s, e = base_window(start_t, end_t)
    s_lo, s_hi = s + timedelta(seconds=min(srange)), s + timedelta(seconds=max(srange))
    e_lo, e_hi = e + timedelta(seconds=min(erange)), e + timedelta(seconds=max(erange))
    rng = lambda a, b: fmt_clock(a) if a == b else f"{fmt_clock(a)} – {fmt_clock(b)}"
    md('<div class="fi-prev">'
       f'<div><div class="k">First entry around</div><div class="v">{rng(s_lo, s_hi)}</div></div>'
       '<div class="arrow">→</div>'
       f'<div><div class="k">Last entry by</div><div class="v">{rng(e_lo, e_hi)}</div></div>'
       '</div>')
    if datetime.now(IST) > s_hi: st.caption("Start time passed, sending begins immediately.")

def render_execution_table(rows):
    if not rows:
        st.caption("No entries processed yet.")
        return
    table_rows = []
    sr_no = 1
    for r in rows:
        kind = r["kind"]
        if kind == "note": continue
        if kind == "ok":
            status = "Successful"
            next_time = f"{r.get('gap_sec', 0)}s"
        elif kind == "final":
            status = "Successful"
            next_time = "Completed"
        elif kind == "fail":
            status = "Failed"
            next_time = "N/A"
        else:
            status = kind.capitalize()
            next_time = "N/A"
        table_rows.append([sr_no, r.get("row", "-"), status, next_time])
        sr_no += 1
    if table_rows:
        html_table(["SR No.", "Row No.", "Status", "Next Entry (s)"], table_rows, ["mono", "mono", "", "mono"])

def view_configure():
    saved = st.session_state.get("inp", {})
    section_head("Setup Your Run", "Complete these four steps. Nothing is sent until you review everything.")
    
    with st.container(border=True):
        step_head("01", "Your Data", "Provide your spreadsheet with one row per response.")
        kinds = ["Upload a file", "Google Sheet link", "Enter data manually"] + (["File on this computer"] if ALLOW_LOCAL_PATHS else [])
        kind_keys = {"Upload a file": "upload", "Google Sheet link": "sheet", "Enter data manually": "manual", "File on this computer": "local"}
        prev = next((k for k, v in kind_keys.items() if v == saved.get("source_kind")), kinds[0])
        kind_label = st.radio("Data Source", kinds, horizontal=True, key="source_kind", index=kinds.index(prev) if prev in kinds else 0, label_visibility="collapsed")
        kind = kind_keys[kind_label]
        upload, sheet_url, local_path, manual_data = None, "", "", ""
        
        if kind == "upload":
            upload = st.file_uploader("Choose a spreadsheet file", type=[e.lstrip(".") for e in SUPPORTED_EXT], key="upload", help="Supported: Excel, CSV, TSV, OpenDocument, JSON")
            if upload: st.success(f"File selected: {upload.name}")
            elif saved.get("upload"): st.info(f"Using previous file: {saved['upload'][0]}")
            st.caption("Tip: Include a 'Status' column to track sent rows.")
        elif kind == "sheet":
            sheet_url = st.text_input("Google Sheet Link", value=saved.get("sheet_url", ""), key="sheet_url", placeholder="https://docs.google.com/spreadsheets/d/...", help="Sheet must be shared as 'Anyone with link'")
            st.caption("Tip: Make sure your Google Sheet is publicly accessible.")
        elif kind == "manual":
            manual_data = st.text_area("Paste CSV Data", value=saved.get("manual_data", ""), key="manual_data", height=200, placeholder="Name,Email,Age\nJohn,john@example.com,30", help="Include headers in first row")
            st.caption("Tip: First line must contain column headers.")
        else:
            local_path = st.text_input("Full File Path", value=saved.get("local_path", ""), key="local_path", placeholder="C:\\data\\responses.xlsx")
            st.caption("Only available if FORMINTELLECT_ALLOW_LOCAL_PATHS=1")

        df_temp = None
        if kind == "upload" and upload:
            try: df_temp = pd.read_csv(io.BytesIO(upload.getvalue())) if upload.name.lower().endswith('.csv') else pd.read_excel(io.BytesIO(upload.getvalue()))
            except: pass
        elif kind == "manual" and manual_data:
            try: df_temp = pd.read_csv(io.StringIO(manual_data))
            except: pass
        
        if df_temp is not None and not df_temp.empty:
            total_rows = len(df_temp)
            min_time_sec = total_rows * MIN_SECONDS_PER_ROW
            md(f'''
            <div class="fi-calc">
                <div><div class="k">Total Rows Detected</div><div class="v">{total_rows}</div><div class="sub">data rows in your file</div></div>
                <div><div class="k">Min Execution Time</div><div class="v">{format_human(min_time_sec)}</div><div class="sub">at {MIN_SECONDS_PER_ROW}s per row minimum</div></div>
            </div>''')

    with st.container(border=True):
        step_head("02", "Your Google Form", "Paste the pre-filled link to map columns to questions.")
        link = st.text_input("Pre-filled Form Link", value=saved.get("link", ""), key="link", placeholder="https://docs.google.com/forms/d/e/.../viewform?entry.123=...", help="This tells the app which column goes to which question")
        with st.expander("How to get a pre-filled link"):
            st.markdown("""**Follow these steps:**\n\n1. Open your form in Google Forms (edit mode)\n2. Click the **⋮ menu** (three dots) at the top right\n3. Select **"Get pre-filled link"**\n4. Type any answer in every question\n5. Click **"Get link"** then **"Copy link"**\n6. Paste the link above\n\n**Important:** Your spreadsheet columns should match the order of questions in the form, or be named like `entry.123456789`.""")

    with st.container(border=True):
        step_head("03", "Which Rows", "Rows marked 'submitted' are automatically skipped.")
        use_random = st.toggle("Pick rows at random", value=bool(saved.get("random", False)), key="random", help="Off: top to bottom. On: random order.")
        sample_n = 0
        if use_random: sample_n = st.number_input("How many rows to send?", min_value=0, step=1, value=int(saved.get("sample_n", 0)), key="sample_n", help="Enter 0 to send all pending rows")

    with st.container(border=True):
        step_head("04", "When to Submit", "Entries spread randomly between these times (IST).")
        
        # Calculate current IST time for dynamic defaults
        now_ist = datetime.now(IST)
        current_time = now_ist.time()
        default_end_time = (now_ist + timedelta(hours=1)).time()

        # Use the custom 12-hour AM/PM picker in separate single cells
        c1, c2 = st.columns(2)
        with c1:
            start = get_time_12h("Start Time (HH:MM:SS AM/PM)", saved.get("start", current_time), "start_time")
        with c2:
            end = get_time_12h("End Time (HH:MM:SS AM/PM)", saved.get("end", default_end_time), "end_time")
        
        st.markdown("---")
        
        with st.container(border=True):
            st.markdown("### Start Delay")
            st.caption("Random delay added before the first entry starts.")
            sc1, sc2 = st.columns(2)
            s_min = saved.get("srange", (0, 0))[0] if isinstance(saved.get("srange"), tuple) else 0
            s_max = saved.get("srange", (0, 0))[1] if isinstance(saved.get("srange"), tuple) else 0
            srange_min = sc1.number_input("Min extra wait (seconds)", min_value=0, max_value=3600, value=int(s_min), step=5, key="srange_min")
            srange_max = sc2.number_input("Max extra wait (seconds)", min_value=0, max_value=3600, value=int(s_max), step=5, key="srange_max")
            
        with st.container(border=True):
            st.markdown("### End Delay")
            st.caption("Random delay added after the scheduled end time.")
            ec1, ec2 = st.columns(2)
            e_min = saved.get("erange", (0, 0))[0] if isinstance(saved.get("erange"), tuple) else 0
            e_max = saved.get("erange", (0, 0))[1] if isinstance(saved.get("erange"), tuple) else 0
            erange_min = ec1.number_input("Min extra time after end (seconds)", min_value=0, max_value=3600, value=int(e_min), step=5, key="erange_min")
            erange_max = ec2.number_input("Max extra time after end (seconds)", min_value=0, max_value=3600, value=int(e_max), step=5, key="erange_max")
            
        srange = (srange_min, srange_max)
        erange = (erange_min, erange_max)
        schedule_preview(start, end, srange, erange)

    if st.button("Check Everything", type="primary", width="stretch"):
        inp = dict(source_kind=kind, upload=(upload.name, upload.getvalue()) if upload else saved.get("upload"), sheet_url=sheet_url, local_path=local_path, manual_data=manual_data, link=link, random=bool(use_random), sample_n=int(sample_n), start=start, end=end, srange=tuple(srange), erange=tuple(erange))
        st.session_state.inp = inp
        try:
            with st.spinner("Reading data and checking form..."):
                st.session_state.prep = prepare(inp, st.session_state.workdir)
            st.rerun()
        except AppError as exc: show_error(exc)
        except Exception as exc: error_card(**explain_exception(exc), where="While checking setup")

def build_checks(prep: Prepared):
    s, cfg = prep.summary, prep.time_cfg
    checks = []
    src = f'{s["source"]} · {s["rows"]} rows · {s["cols"]} columns · {s["pending"]} pending'
    checks.append(("ok", "Data loaded", src))
    checks.append(("ok", "Rows to send", f'{s["selected"]} of {s["pending"]} rows' + (", random order" if s["random"] else ", top to bottom")))
    checks.append(("ok", "Columns matched", f'{len(prep.mapping)} columns matched to {s["entry_fields"]} questions'))
    if prep.form.field_types: checks.append(("ok", "Form read", "Question types and choices loaded"))
    elif s["form_reachable"]: checks.append(("warn", "Form partly read", "Questions couldn't be fully read"))
    else: checks.append(("warn", "Couldn't open form", "Answers couldn't be validated"))
    if prep.corrections: checks.append(("warn", "Auto-fixes applied", f"{len(prep.corrections)} values adjusted"))
    else: checks.append(("ok", "No fixes needed", "All values look good"))
    checks.append(("ok", "Validation passed", f'{s["selected"]} rows × {s["validated_fields"]} questions checked'))
    sched = f'{fmt_clock(prep.rand_start)} → {fmt_clock(prep.rand_end)} IST'
    if datetime.now(IST) >= prep.rand_start: checks.append(("warn", "Schedule", sched + " (start time passed)"))
    else: checks.append(("ok", "Schedule", sched))
    return checks

def view_preflight(prep: Prepared):
    section_head("Review Before Sending", "Everything looks good. Nothing has been sent yet.")
    cfg = prep.time_cfg
    callout("ok", f"<b>Ready to send {cfg['total_rows']} row(s)</b> between <b>{fmt_clock(prep.rand_start)}</b> and <b>{fmt_clock(prep.rand_end)}</b> IST.")
    icons = {"ok": "✓", "warn": "!", "err": "✕"}
    rows = "".join(f'<div class="fi-check {k}"><span class="ic">{icons[k]}</span><div><div class="t">{esc(t)}</div><div class="d">{esc(d)}</div></div></div>' for k, t, d in build_checks(prep))
    md(f'<div class="fi-checks">{rows}</div>')
    t_map, t_fix, t_sched, t_console = st.tabs(["Column Mapping", "Auto-fixes", "Schedule", "Full Log"])
    with t_map: html_table(["Column", "Question", "ID", "Type", "Required"], [[m["Column"], m["Question"], m["ID"], m["Type"], m["Required"]] for m in prep.mapping], ["", "", "mono", "mono", ""])
    with t_fix:
        if prep.corrections:
            st.caption("Values adjusted to match form exactly.")
            html_table(["Row", "Column", "Was", "Now", "Why"], [[c["Row"], c["Column"], c["From"], c["To"], c["Why"]] for c in prep.corrections], ["mono", "", "val", "fix", ""])
        else: st.caption("No values needed changing.")
    with t_sched:
        es, ee = cfg["entered_start"], cfg["entered_end"]
        stats_grid(stat("You entered", fmt_clock(es), es.strftime("%d %b %Y"), small=True), stat("Until", fmt_clock(ee), ee.strftime("%d %b %Y"), small=True), stat("Actual start", fmt_clock(cfg["rand_start"]), f'+{cfg["rs"]}s delay', "g", small=True), stat("Actual end", fmt_clock(cfg["rand_end"]), f'+{cfg["re"]}s extra', "g", small=True))
        avg = cfg["window"] / max(cfg["total_rows"], 1)
        stats_grid(stat("Window entered", format_human(cfg["original_window"]), "", small=True), stat("Window used", format_human(cfg["window"]), "", "g", small=True), stat("Rows to send", cfg["total_rows"], f'needs {format_human(cfg["min_required"])}', small=True), stat("Average gap", format_human(avg), "between entries", small=True))
        if st.button("Pick new random times"):
            try: reroll_schedule(prep); st.rerun()
            except AppError as exc: show_error(exc)
    with t_console: st.code(console_text(prep), language=None)
    md("<div style='height:1rem'></div>")
    b1, b2 = st.columns([1, 2])
    if b1.button("Back to Setup", width="stretch"): st.session_state.prep = None; st.rerun()
    if b2.button("Start Sending", type="primary", width="stretch"): job = Job(prep); job.start(); st.session_state.job = job; st.rerun()

def render_failures(failures):
    html_table(["Row", "Time", "Problem", "Fix"], [[f["row"], f["time"], f["problem"], f["fix"]] for f in failures], ["row", "mono", "", "fix"])

def render_monitor(job: Job):
    s = job.snapshot()
    label, tone, sub = STATUS_LABEL[s["status"]]
    total = max(s["total"], 1)
    ok_pct = s["submitted"] / total * 100
    bad_pct = s["failed"] / total * 100
    if s["next_target"] is not None and s["status"] in ("waiting", "running"):
        secs = max(0, int((s["next_target"] - datetime.now(IST)).total_seconds()))
        nxt_v, nxt_s = format_mm_ss(secs), f'at {fmt_clock(s["next_target"])} IST'
    else: nxt_v, nxt_s = "—", ""
    md('<div class="fi-stats">' + stat("Status", label, sub, tone) + stat("Sent", f'{s["submitted"]}/{s["total"]}', f"{int(ok_pct)}% done", "g", raw_value=True) + stat("Failed", s["failed"], "see below" if s["failed"] else "none", "r" if s["failed"] else "") + stat("Next in", nxt_v, nxt_s) + f'</div><div class="fi-bar"><div class="ok" style="width:{ok_pct:.1f}%"></div><div class="bad" style="width:{bad_pct:.1f}%"></div></div>')
    if s["failures"]:
        callout("err", f"<b>{len(s['failures'])} row(s) failed.</b> Failed rows stay pending.")
        render_failures(s["failures"])
    
    md('<div class="fi-tbl-h"><span>Execution Log</span><span>{0} entries</span></div>'.format(len([r for r in s["rows"] if r["kind"] != "note"])))
    render_execution_table(s["rows"])

    if job.active:
        if st.button("Stop Sending", key="stop_btn"): job.stop.set()
    elif not job.notified: job.notified = True; st.rerun()

def view_execute(job: Job):
    section_head("Sending in Progress", "Keep this tab open until complete.")
    monitor = st.fragment(run_every=1)(render_monitor)
    monitor(job)
    with st.expander("View Raw Log", expanded=False): st.code(console_text(job.prep), language=None)

def view_results(job: Job):
    s = job.snapshot()
    section_head("Results", "Summary of what was sent.")
    if s["status"] == "error":
        err = s["error"] or {}
        error_card(err.get("title", "Error"), "During sending", err.get("why", ""), err.get("fix", ""), technical=err.get("technical", ""))
    elif s["status"] == "stopped": callout("warn", "<b>Stopped by user.</b> Sent rows marked 'submitted'.")
    elif s["failed"]: callout("warn", f"<b>{s['submitted']} sent, {s['failed']} failed.</b> Failed rows stay pending.")
    else: callout("ok", f"<b>All {s['submitted']} row(s) sent successfully.</b>")
    stats_grid(stat("Sent", f'{s["submitted"]}/{s["total"]}', "", "g", raw_value=True), stat("Failed", s["failed"], "", "r" if s["failed"] else ""), stat("Status", STATUS_LABEL[s["status"]][0], STATUS_LABEL[s["status"]][2], STATUS_LABEL[s["status"]][1]), stat("Run time", format_human(job.elapsed()), "total"))
    if s["failures"]:
        md('<div class="fi-tbl-h"><span>Failed Rows</span></div>')
        render_failures(s["failures"])
        with st.expander("Technical Details"): st.code("\n\n".join(f"Row {f['row']}:\n{f['technical']}" for f in s["failures"]), language=None)
    
    md('<div class="fi-tbl-h"><span>Final Execution Log</span><span>{0} entries</span></div>'.format(len([r for r in s["rows"] if r["kind"] != "note"])))
    render_execution_table(s["rows"])

    md("<div style='height:1rem'></div>")
    cols = st.columns(4 if s["failures"] else 3)
    out = job.prep.output_path
    if os.path.exists(out):
        with open(out, "rb") as fh: cols[0].download_button("Download Updated File", fh.read(), file_name=os.path.basename(out), width="stretch")
    else: cols[0].button("No File", disabled=True, width="stretch")
    cols[1].download_button("Download Log", job.log_text(), file_name="formintellect_log.txt", width="stretch")
    if s["failures"]: cols[2].download_button("Download Failed Rows", job.failures_csv(), file_name="failed_rows.csv", width="stretch")
    if cols[-1].button("Start New Run", type="primary", width="stretch"): st.session_state.pop("job", None); st.session_state.pop("prep", None); st.rerun()
    if s["status"] == "done" and out.lower().endswith(".xlsx"): st.caption("Updated file includes 'Analysis' sheet.")
    with st.expander("View Raw Log", expanded=False): st.code(console_text(job.prep), language=None)

# ============================================================
#  MAIN APP
# ============================================================

def current_stage():
    job = st.session_state.get("job")
    if job: return "execute" if job.active else "results"
    return "preflight" if st.session_state.get("prep") else "configure"

def main():
    st.set_page_config(page_title="FormIntellect", page_icon="✅", layout="centered")
    inject_css(CSS)
    if "workdir" not in st.session_state: st.session_state.workdir = tempfile.mkdtemp(prefix="formintellect_")
    
    md('<div class="fi-hero"><div class="fi-logo">✓</div><div class="fi-name">Form<span>Intellect</span></div></div><div class="fi-tag">Professional Google Form submission scheduler</div>')
    stage = current_stage()
    stepper(stage)
    if stage == "configure": view_configure()
    elif stage == "preflight": view_preflight(st.session_state.prep)
    elif stage == "execute": view_execute(st.session_state.job)
    else: view_results(st.session_state.job)

if __name__ == "__main__":
    main()
